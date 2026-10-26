import math as math_, random, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

NAVY=RGBColor(0x2E,0x5E,0x8E); TEAL=RGBColor(0x1E,0x8A,0x82); RED=RGBColor(0xC0,0x39,0x2B)
TXT=RGBColor(0x1F,0x3A,0x5C); GRAY=RGBColor(0x6B,0x72,0x80); SLATE=RGBColor(0x2B,0x3A,0x4A)
WHITE=RGBColor(0xFF,0xFF,0xFF); LNAVY=RGBColor(0xEE,0xF3,0xF9); LTEAL=RGBColor(0xE8,0xF5,0xF3)
LGRAY=RGBColor(0x9A,0xA5,0xB4); NEUR=RGBColor(0xD6,0xE4,0xF2)
FONT="Noto Sans CJK JP"
import os, matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
plt.rcParams["mathtext.fontset"]="cm"
MATHDIR=os.path.join(os.path.dirname(os.path.abspath(sys.argv[1])),"_math"); os.makedirs(MATHDIR,exist_ok=True)
_mc=[0]
def mathimg(tex,size,color):
    _mc[0]+=1; p=os.path.join(MATHDIR,f"m{_mc[0]}.png")
    fig=plt.figure(figsize=(0.01,0.01)); fig.text(0,0,tex,fontsize=size,color="#%02X%02X%02X"%tuple(color))
    fig.savefig(p,dpi=600,transparent=True,bbox_inches="tight",pad_inches=0.02); plt.close(fig)
    from PIL import Image
    w,h=Image.open(p).size; return p,w/600,h/600

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6])
bg=s.background.fill; bg.solid(); bg.fore_color.rgb=WHITE

def text(shape, lines, size=14, color=TXT, bold=False, align=PP_ALIGN.CENTER):
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(0.05); tf.margin_top=tf.margin_bottom=Inches(0.02)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    if isinstance(lines,str): lines=[(lines,size,color,bold)]
    for i,(t,sz,c,b) in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=t; f=r.font; f.size=Pt(sz); f.color.rgb=c; f.bold=b; f.name=FONT
        rPr=r._r.get_or_add_rPr(); ea=rPr.find(qn('a:ea'))
        if ea is None: ea=etree.SubElement(rPr,qn('a:ea'))
        ea.set('typeface',FONT)

def box(x,y,w,h,fill=WHITE,line=NAVY,lw=1.4,shape=MSO_SHAPE.ROUNDED_RECTANGLE,adj=0.16):
    sh=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=adj
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False
    return sh

def label(x,y,w,h,lines,size=12,color=GRAY,bold=False,align=PP_ALIGN.CENTER):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); text(tb,lines,size,color,bold,align); return tb

def line(x1,y1,x2,y2,color=LGRAY,w=1.0,arrow=True,dash=False):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w)
    ln=c.line._get_or_add_ln()
    if dash:
        d=etree.SubElement(ln,qn('a:prstDash')); d.set('val','dash')
    if arrow:
        t=etree.SubElement(ln,qn('a:tailEnd')); t.set('type','triangle'); t.set('w','lg'); t.set('len','lg')
    return c


def zone(x,y,w,h,tab,accent,fill,size=18):
    box(x,y,w,h,fill=fill,line=accent,lw=1.8,adj=0.05)
    tw=size/72*1.05*len(tab)+0.6
    t=box(x+0.25,y-0.27,tw,0.54,fill=accent,line=None,adj=0.5); text(t,tab,size,WHITE,True)

def math(cx,cy,tex,size,color=TXT):
    p,w,h=mathimg(tex,size,color)
    s.shapes.add_picture(p,Inches(cx-w/2),Inches(cy-h/2),Inches(w),Inches(h)); return w

def row(cx,cy,items):
    """items: ('t',str,size,color,bold) or ('m',tex,size,color) を横に並べて中央揃え"""
    rend=[]
    for it in items:
        if it[0]=='m':
            p,w,h=mathimg(it[1],it[2],it[3]); rend.append(('m',p,w,h))
        else:
            w=sum((it[2]/72)*(1.02 if ord(ch)>0x2000 else 0.66) for ch in it[1])+0.2
            rend.append(('t',it,w,it[2]/72*1.6))
    tot=sum(r[2] for r in rend); x=cx-tot/2
    for r in rend:
        if r[0]=='m':
            s.shapes.add_picture(r[1],Inches(x),Inches(cy-r[3]/2),Inches(r[2]),Inches(r[3]))
        else:
            it=r[1]; tb=label(x,cy-r[3]/2,r[2],r[3],it[1],it[2],it[3],it[4]); tb.text_frame.word_wrap=False
        x+=r[2]

ZT,ZH=0.95,4.35   # zones top/height
# zones
zone(0.35,ZT,2.35,ZH,"入力層",NAVY,LNAVY)
zone(3.25,ZT,5.75,ZH,"リザバー",NAVY,LNAVY)
zone(9.55,ZT,3.45,ZH,"読み出し層",TEAL,LTEAL)

R=0.23
inputs=[(1.15,1.85),(1.15,2.95),(1.15,4.05)]
random.seed(7)
cols=[4.0,5.25,6.5,7.75]; rows=[1.75,2.85,3.95]
neur=[]
for r in rows:
    for c in cols:
        neur.append((c+random.uniform(-0.3,0.3), r+random.uniform(-0.2,0.2)))
neur.append((4.6,2.3)); neur.append((7.1,3.4))
def edge(p,q,rp,rq,**kw):
    dx,dy=q[0]-p[0],q[1]-p[1]; d=math_.hypot(dx,dy); ux,uy=dx/d,dy/d
    line(p[0]+ux*rp,p[1]+uy*rp,q[0]-ux*rq,q[1]-uy*rq,**kw)
N=len(neur)
wo=(9.95,2.2,1.7,1.5)
for (x,y) in neur:
    tx=wo[0]; ty=wo[1]+0.15+(y-1.5)/(4.2-1.5)*(wo[3]-0.3)
    edge((x,y),(tx,ty),R,0.0,color=RGBColor(0x9F,0xD3,0xCC),w=0.9,arrow=False)
edges=set()
for i in range(N):
    cand=sorted(range(N),key=lambda j: math_.dist(neur[i],neur[j]))[1:6]
    for j in random.sample(cand,2):
        if (j,i) in edges and random.random()<0.5: continue
        edges.add((i,j))
for i,j in edges: edge(neur[i],neur[j],R,R+0.02,color=NAVY,w=1.0)
for (x,y) in inputs:
    tg=sorted(range(N),key=lambda j: math_.dist((x,y),neur[j]))[:3]
    for j in tg: edge((x,y),neur[j],0.36,R+0.02,color=NAVY,w=1.0)
for (x,y) in neur:
    box(x-R,y-R,2*R,2*R,fill=NEUR,line=NAVY,lw=1.3,shape=MSO_SHAPE.OVAL)
for i,(x,y) in enumerate(inputs):
    box(x-0.36,y-0.36,0.72,0.72,fill=WHITE,line=NAVY,lw=1.8,shape=MSO_SHAPE.OVAL)
    math(x,y,f"$u_{i+1}$",24)

# zone bottom labels
LB=ZT+ZH-0.42
row(1.525,LB,[('t',"入力",18,TXT,True),('m',r"$\boldsymbol{u}(t)$",22,TXT)])
row(6.125,LB,[('t',"状態",18,TXT,True),('m',r"$\boldsymbol{x}(t)$",22,TXT)])
row(12.1,LB,[('t',"出力",18,TXT,True),('m',r"$y(t)$",22,TXT)])

# readout
b=box(*wo,fill=WHITE,line=TEAL,lw=2.0)
math(wo[0]+wo[2]/2,wo[1]+0.55,r"$W_{\mathrm{out}}$",28)
label(wo[0],wo[1]+0.85,wo[2],0.45,"学習する",18,TEAL,True)
yc=(12.4,wo[1]+wo[3]/2)
line(wo[0]+wo[2],yc[1],yc[0]-0.38,yc[1],color=TEAL,w=1.8)
box(yc[0]-0.38,yc[1]-0.38,0.76,0.76,fill=WHITE,line=TEAL,lw=2.0,shape=MSO_SHAPE.OVAL)
math(yc[0],yc[1],r"$y$",26)

# weight labels
row(3.0,0.42,[('m',r"$W_{\mathrm{in}}$",24,SLATE),('t',"（固定）",18,SLATE,True)])

# feedback
fy=0.6
line(yc[0],yc[1]-0.38,yc[0],fy,color=RED,w=1.5,arrow=False,dash=True)
line(yc[0],fy,7.3,fy,color=RED,w=1.5,arrow=False,dash=True)
line(7.3,fy,7.3,ZT,color=RED,w=1.5,arrow=True,dash=True)
row(9.85,0.27,[('t',"出力フィードバック",16,RED,False),('m',r"$W_{\mathrm{fb}}$",20,RED),('t',"（任意・FORCE学習）",16,RED,False)])

# bottom: equations
BT,BH=5.65,1.65
eq=box(3.25,BT,5.75,BH,fill=WHITE,line=NAVY,lw=1.4)
math(6.125,BT+0.38,r"$\boldsymbol{x}(t+1)=\tanh\left(W_{\mathrm{in}}\,\boldsymbol{u}(t+1)+W\,\boldsymbol{x}(t)\right)$",22)
math(6.125,BT+0.88,r"$y(t)=W_{\mathrm{out}}\,\boldsymbol{x}(t)$",22)
row(6.125,BT+1.33,[('m',r"$W$",18,GRAY),('t',"：ランダム・固定（",16,GRAY,False),('m',r"$\rho(W)<1$",18,GRAY),('t',"が目安）",16,GRAY,False)])
rr=box(9.55,BT,1.95,BH,fill=TEAL,line=None)
text(rr,[("リッジ回帰",20,WHITE,True),("一度に解く",16,WHITE,False)])
tg=box(11.75,BT,1.25,BH,fill=WHITE,line=TEAL,lw=1.6)
label(11.75,BT+0.25,1.25,0.45,"目標",18,TXT,True)
math(12.375,BT+1.0,r"$d(t)$",24)
line(11.75,BT+BH/2,11.5,BT+BH/2,color=TEAL,w=1.8)
line(10.525,BT,10.525,wo[1]+wo[3],color=TEAL,w=1.8,dash=True)

# legend
lx,ly=0.4,5.75
for k,(col,dash,lab) in enumerate([(NAVY,False,"固定（学習しない）"),(TEAL,False,"学習する"),(RED,True,"任意の結合")]):
    yy=ly+k*0.5
    line(lx,yy+0.22,lx+0.55,yy+0.22,color=col,w=2.0,arrow=False,dash=dash)
    label(lx+0.62,yy,2.2,0.45,lab,16,TXT,False,PP_ALIGN.LEFT)

prs.save(sys.argv[1])
import shutil; shutil.rmtree(MATHDIR)
