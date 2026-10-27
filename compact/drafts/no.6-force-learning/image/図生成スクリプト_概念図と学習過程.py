# 図1（概念図）と図3（学習の過程）を作る
# 使い方: python3 gen.py <出力ディレクトリ>
import math as math_, os, sys
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["mathtext.fontset"]="cm"

OUT=sys.argv[1]
NAVY=RGBColor(0x2E,0x5E,0x8E); TEAL=RGBColor(0x1E,0x8A,0x82); RED=RGBColor(0xC0,0x39,0x2B)
TXT=RGBColor(0x1F,0x3A,0x5C); GRAY=RGBColor(0x6B,0x72,0x80); SLATE=RGBColor(0x2B,0x3A,0x4A)
WHITE=RGBColor(0xFF,0xFF,0xFF); LNAVY=RGBColor(0xEE,0xF3,0xF9); LTEAL=RGBColor(0xE8,0xF5,0xF3)
LSLATE=RGBColor(0xF1,0xF3,0xF5); LGRAY=RGBColor(0x9A,0xA5,0xB4)
FONT="Noto Sans CJK JP"
hexc=lambda c:"#%02X%02X%02X"%tuple(c)
TMP=os.path.join(OUT,"_tmp"); os.makedirs(TMP,exist_ok=True)
_mc=[0]
def tmpfile():
    _mc[0]+=1; return os.path.join(TMP,f"i{_mc[0]}.png")
def mathimg(tex,size,color):
    p=tmpfile()
    fig=plt.figure(figsize=(0.01,0.01)); fig.text(0,0,tex,fontsize=size,color=hexc(color))
    fig.savefig(p,dpi=600,transparent=True,bbox_inches="tight",pad_inches=0.02); plt.close(fig)
    from PIL import Image
    w,h=Image.open(p).size; return p,w/600,h/600

s=None
def new_prs():
    global s
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    s=prs.slides.add_slide(prs.slide_layouts[6])
    bg=s.background.fill; bg.solid(); bg.fore_color.rgb=WHITE
    return prs

def text(shape, lines, size=14, color=TXT, bold=False, align=PP_ALIGN.CENTER):
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(0.05); tf.margin_top=tf.margin_bottom=Inches(0.02)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    if isinstance(lines,str): lines=[(lines,size,color,bold)]
    for i,(t,sz,c,b) in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=t; f=r.font; f.size=Pt(sz); f.color.rgb=c; f.bold=b; f.name=FONT
        rPr=r._r.get_or_add_rPr(); ea=rPr.find(qn('a:ea'))
        if ea is None: ea=etree.SubElement(rPr,qn('a:ea'))
        ea.set('typeface',FONT)

def box(x,y,w,h,fill=WHITE,line=NAVY,lw=1.4,shape=MSO_SHAPE.ROUNDED_RECTANGLE,adj=0.16):
    sh=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=adj
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False
    return sh

def label(x,y,w,h,lines,size=12,color=GRAY,bold=False,align=PP_ALIGN.CENTER):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); text(tb,lines,size,color,bold,align); return tb

def line(x1,y1,x2,y2,color=LGRAY,w=1.0,arrow=True,dash=False):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w)
    ln=c.line._get_or_add_ln()
    if dash:
        d=etree.SubElement(ln,qn('a:prstDash')); d.set('val','dash')
    if arrow:
        t=etree.SubElement(ln,qn('a:tailEnd')); t.set('type','triangle'); t.set('w','lg'); t.set('len','lg')
    return c

def zone(x,y,w,h,tab,accent,fill,size=18):
    box(x,y,w,h,fill=fill,line=accent,lw=1.8,adj=0.05)
    tw=size/72*1.05*len(tab)+0.6
    t=box(x+0.25,y-0.27,tw,0.54,fill=accent,line=None,adj=0.5); text(t,tab,size,WHITE,True)

def math(cx,cy,tex,size,color=TXT):
    p,w,h=mathimg(tex,size,color)
    s.shapes.add_picture(p,Inches(cx-w/2),Inches(cy-h/2),Inches(w),Inches(h)); return w

def row(cx,cy,items):
    """items: ('t',str,size,color,bold) or ('m',tex,size,color) を横に並べて中央揃え"""
    rend=[]
    for it in items:
        if it[0]=='m':
            p,w,h=mathimg(it[1],it[2],it[3]); rend.append(('m',p,w,h))
        else:
            w=sum((it[2]/72)*(1.02 if ord(ch)>0x2000 else 0.66) for ch in it[1])+0.2
            rend.append(('t',it,w,it[2]/72*1.6))
    tot=sum(r[2] for r in rend); x=cx-tot/2
    for r in rend:
        if r[0]=='m':
            s.shapes.add_picture(r[1],Inches(x),Inches(cy-r[3]/2),Inches(r[2]),Inches(r[3]))
        else:
            it=r[1]; tb=label(x,cy-r[3]/2,r[2],r[3],it[1],it[2],it[3],it[4]); tb.text_frame.word_wrap=False
        x+=r[2]

# ---------------- 図1：概念図 ----------------
prs=new_prs()
zone(0.35,1.15,9.75,3.05,"ネットワーク",NAVY,LNAVY)
zone(0.35,4.85,12.6,2.3,"学習（学習中だけ動かす）",TEAL,LTEAL)

# リザバー
rv=(0.85,1.75,3.5,1.95)
box(*rv,fill=WHITE,line=NAVY,lw=1.6)
label(rv[0],rv[1]+0.18,rv[2],0.5,"リザバー",20,TXT,True)
row(rv[0]+rv[2]/2,rv[1]+0.95,[('t',"再帰結合",15,GRAY,False),('m',r"$J$",18,GRAY),('t',"：ランダム・固定",15,GRAY,False)])
label(rv[0],rv[1]+1.3,rv[2],0.45,"カオス的に動く",16,GRAY,False)
# 読み出し
wo=(5.9,1.95,2.0,1.55)
box(*wo,fill=WHITE,line=TEAL,lw=2.0)
row(wo[0]+wo[2]/2,wo[1]+0.5,[('t',"読み出し",18,TXT,True),('m',r"$\boldsymbol{w}$",24,TXT)])
label(wo[0],wo[1]+0.85,wo[2],0.45,"学習する",18,TEAL,True)
cy=wo[1]+wo[3]/2
line(rv[0]+rv[2],cy,wo[0],cy,color=NAVY,w=1.8)
row((rv[0]+rv[2]+wo[0])/2,cy-0.35,[('t',"状態",16,TXT,True),('m',r"$\boldsymbol{r}(t)$",20,TXT)])
# 出力
zc=(8.8,cy); ZR=0.45
line(wo[0]+wo[2],cy,zc[0]-ZR,cy,color=TEAL,w=1.8)
box(zc[0]-ZR,cy-ZR,2*ZR,2*ZR,fill=WHITE,line=TEAL,lw=2.0,shape=MSO_SHAPE.OVAL)
math(zc[0],cy,r"$z(t)$",22)
label(zc[0]+ZR+0.02,cy-0.25,0.85,0.5,"出力",17,TXT,True)
# フィードバック（上を回る）
fy=0.55; fbx=rv[0]+rv[2]*0.72
line(zc[0],cy-ZR,zc[0],fy,color=RED,w=1.8,arrow=False)
line(zc[0],fy,fbx,fy,color=RED,w=1.8,arrow=False)
line(fbx,fy,fbx,rv[1],color=RED,w=1.8)
row((fbx+zc[0])/2,fy-0.3,[('t',"出力フィードバック",16,RED,True),('m',r"$\boldsymbol{w}_{\mathrm{f}}$",20,RED),('t',"（固定）",16,RED,True)])

# 学習ループ
LY=6.05
rls=(5.7,LY-0.6,2.4,1.2)
b=box(*rls,fill=TEAL,line=None); text(b,[("RLS",22,WHITE,True),("毎ステップ直す",16,WHITE,False)])
ec=(zc[0],LY); ER=0.55
box(ec[0]-ER,LY-ER,2*ER,2*ER,fill=WHITE,line=TEAL,lw=2.0,shape=MSO_SHAPE.OVAL)
math(ec[0],LY-0.13,r"$e(t)$",22)
label(ec[0]-ER,LY+0.1,2*ER,0.35,"誤差",15,TEAL,True)
tg=(10.6,LY-0.6,1.9,1.2)
box(*tg,fill=WHITE,line=TEAL,lw=1.6)
label(tg[0],tg[1]+0.12,tg[2],0.45,"目標",17,TXT,True)
math(tg[0]+tg[2]/2,tg[1]+0.82,r"$f(t)$",22)
line(zc[0],cy+ZR,ec[0],LY-ER,color=TEAL,w=1.8)
line(tg[0],LY,ec[0]+ER,LY,color=TEAL,w=1.8)
line(ec[0]-ER,LY,rls[0]+rls[2],LY,color=TEAL,w=1.8)
line(rls[0]+rls[2]/2,rls[1],wo[0]+wo[2]/2,wo[1]+wo[3],color=TEAL,w=1.8,dash=True)
math((ec[0]+ER+tg[0])/2,LY-0.35,r"$e=z-f$",18,GRAY)
label(7.0,4.3,2.4,0.4,"重みを直す",15,TEAL,True,PP_ALIGN.LEFT)

# 凡例
lx,ly=10.5,1.5
for k,(col,lab) in enumerate([(NAVY,"固定（学習しない）"),(TEAL,"学習する"),(RED,"フィードバック")]):
    yy=ly+k*0.55
    line(lx,yy+0.22,lx+0.5,yy+0.22,color=col,w=2.0,arrow=False)
    tb=label(lx+0.55,yy,2.5,0.45,lab,15,TXT,False,PP_ALIGN.LEFT); tb.text_frame.word_wrap=False
prs.save(os.path.join(OUT,"01_FORCE学習の概念図.pptx"))

# ---------------- 図3：学習の過程 ----------------
def simulate(seed=0):
    rng=np.random.default_rng(seed)
    N,g,dt,tau=500,1.5,0.1,1.0
    T1,T2,T3=1500,3000,1500
    J=g*rng.standard_normal((N,N))/np.sqrt(N); wf=rng.uniform(-1,1,N)
    w=0.3*rng.standard_normal(N)/np.sqrt(N); P=np.eye(N)
    n=T1+T2+T3; t=np.arange(n)*dt
    f=np.sin(t/4)+0.5*np.sin(t/2)
    x=0.5*rng.standard_normal(N); r=np.tanh(x); z=w@r
    for _ in range(1000):
        x+=dt/tau*(-x+J@r+wf*z); r=np.tanh(x); z=w@r
    Z=np.zeros(n); R=np.zeros((n,3)); DW=np.full(n,np.nan)
    for i in range(n):
        x+=dt/tau*(-x+J@r+wf*z); r=np.tanh(x); z=w@r
        if T1<=i<T1+T2:
            e=z-f[i]; Pr=P@r; P-=np.outer(Pr,Pr)/(1+r@Pr); dw=-e*(P@r); w+=dw; DW[i]=np.linalg.norm(dw)
        Z[i]=z; R[i]=r[:3]
    return t,f,Z,R,DW,(T1,T2,T3)
t,f,Z,R,DW,(T1,T2,T3)=simulate()
print("RMSE after:",np.sqrt(np.mean((Z[T1+T2:]-f[T1+T2:])**2)))
seg=[slice(0,T1),slice(T1,T1+T2),slice(T1+T2,T1+T2+T3)]
k=50; DWs=np.full_like(DW,np.nan)
m=np.convolve(DW[T1:T1+T2],np.ones(k)/k,mode="same"); DWs[T1:T1+T2]=m
DWs[T1:T1+k//2]=np.nan; DWs[T1+T2-k//2:T1+T2]=np.nan

def plot(kind,sl,w_in,h_in):
    p=tmpfile()
    fig=plt.figure(figsize=(w_in,h_in)); ax=fig.add_axes([0,0,1,1]); ax.axis("off")
    tt=t[sl]
    if kind=="z":
        ax.axhline(0,color="#D5DAE1",lw=1.0)
        ax.plot(tt,Z[sl],color=hexc(TEAL),lw=3.2)
        ax.plot(tt,f[sl],color=hexc(SLATE),lw=1.6,ls=(0,(3,2)))
        ax.set_ylim(-1.6,1.6)
    elif kind=="r":
        ax.axhline(0,color="#D5DAE1",lw=1.0)
        for j,c in enumerate(["#2E5E8E","#8FB3D9","#1F3A5C"]):
            ax.plot(tt,R[sl,j],color=c,lw=2.0)
        ax.set_ylim(-1.1,1.1)
    else:
        ax.semilogy(tt,DWs[sl],color=hexc(TEAL),lw=2.4)
        ax.set_ylim(5e-5,2e-2)
    ax.set_xlim(tt[0],tt[-1])
    fig.savefig(p,dpi=300,transparent=True); plt.close(fig); return p

prs=new_prs()
X0,XW=2.25,10.75; GAP=0.25
tot=T1+T2+T3; ws=[(XW-2*GAP)*T/tot for T in (T1,T2,T3)]
xs=[X0, X0+ws[0]+GAP, X0+ws[0]+ws[1]+2*GAP]
ZT,ZB=0.85,6.3
tabs=[("① 学習前",SLATE,LSLATE),("② 学習中",TEAL,LTEAL),("③ 学習後",NAVY,LNAVY)]
rows_y=[(1.4,1.55),(3.2,1.4),(4.85,1.1)]   # (上端, 高さ)
PAD=0.12
for i in range(3):
    zone(xs[i],ZT,ws[i],ZB-ZT,tabs[i][0],tabs[i][1],tabs[i][2])
    for kind,(yy,hh) in zip(["z","r","w"],rows_y):
        if kind=="w" and i!=1:
            label(xs[i],yy,ws[i],hh,"学習しない",15,GRAY,False); continue
        p=plot(kind,seg[i],ws[i]-2*PAD,hh)
        s.shapes.add_picture(p,Inches(xs[i]+PAD),Inches(yy),Inches(ws[i]-2*PAD),Inches(hh))
yy,hh=rows_y[2]
label(xs[1]+0.1,yy-0.1,0.4,0.35,"大",12,GRAY,False); label(xs[1]+0.1,yy+hh-0.25,0.4,0.35,"小",12,GRAY,False)
# 各フェーズの説明
cap=[["出力はでたらめ","（カオス的）"],["すぐ目標に重なる","重みの変化はだんだん小さくなる"],["RLSを止めても","目標を出し続ける"]]
for i in range(3):
    label(xs[i],ZB+0.05,ws[i],0.8,[(cap[i][0],15,TXT,True),(cap[i][1],14,GRAY,False)])
# 行ラベル（左）
def rowlabel(yc,title,tex):
    label(0.1,yc-0.5,2.05,0.45,title,15,TXT,True)
    math(1.125,yc+0.15,tex,20)
rowlabel(rows_y[0][0]+rows_y[0][1]/2,"出力",r"$z(t)$")
rowlabel(rows_y[1][0]+rows_y[1][1]/2,"ニューロン",r"$r_i(t)$")
rowlabel(rows_y[2][0]+rows_y[2][1]/2,"重みの変化",r"$\|\Delta\boldsymbol{w}\|$")
# 凡例（出力の行）
lx,ly=2.3,0.0
line(lx,ly+0.22,lx+0.5,ly+0.22,color=TEAL,w=3.0,arrow=False)
row(lx+1.25,ly+0.22,[('t',"出力",14,TXT,False),('m',r"$z(t)$",16,TXT)])
c=line(lx+2.3,ly+0.22,lx+2.8,ly+0.22,color=SLATE,w=1.6,arrow=False,dash=True)
row(lx+3.55,ly+0.22,[('t',"目標",14,TXT,False),('m',r"$f(t)$",16,TXT)])
# 時間の矢印
line(X0,7.25,X0+XW,7.25,color=LGRAY,w=1.5)
label(1.2,7.05,1.0,0.4,"時間",14,GRAY,False)
prs.save(os.path.join(OUT,"03_FORCE学習の学習過程.pptx"))
import shutil; shutil.rmtree(TMP)
