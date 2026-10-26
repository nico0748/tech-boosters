# compact no.3 ナブラ演算子の図（勾配・発散・回転）を生成するスクリプト
# 使い方: python3 図生成スクリプト.py <出力ディレクトリ> <作業ディレクトリ>
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = sys.argv[1]
WORK = sys.argv[2]
os.makedirs(WORK, exist_ok=True)

NAVY = "#2E5E8E"
TEAL = "#1E8A82"
TEXT = "#1F3A5C"
GRAY = "#6B7280"
LIGHT = {NAVY: "#EEF3F9", TEAL: "#E9F5F3"}
FONT = "Noto Sans CJK JP"

plt.rcParams["mathtext.fontset"] = "cm"
plt.rcParams["font.family"] = FONT

_count = [0]


def tmp(name):
    _count[0] += 1
    return os.path.join(WORK, f"{_count[0]:02d}_{name}.png")


def rgb(hex_):
    return RGBColor.from_string(hex_.lstrip("#"))


# ---------- matplotlib ----------

def mathimg(tex, size_pt=22, color=TEXT):
    """TeX 風の式を透明 PNG にして (path, 幅in, 高さin) を返す"""
    path = tmp("math")
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${tex}$", fontsize=size_pt, color=color)
    fig.savefig(path, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)
    from PIL import Image
    w, h = Image.open(path).size
    return path, w / 300, h / 300


LIM = 2.2
GRID = np.linspace(-1.8, 1.8, 7)


def new_panel(size_in):
    fig, ax = plt.subplots(figsize=(size_in, size_in))
    ax.set_xlim(-LIM, LIM)
    ax.set_ylim(-LIM, LIM)
    ax.set_aspect("equal")
    ax.set_xticks([-2, 0, 2])
    ax.set_yticks([-2, 0, 2])
    ax.tick_params(labelsize=12, colors=GRAY, length=3)
    for s in ax.spines.values():
        s.set_color("#C9CFD8")
    ax.set_xlabel("$x$", fontsize=14, color=TEXT, labelpad=0)
    ax.set_ylabel("$y$", fontsize=14, color=TEXT, labelpad=0, rotation=0)
    ax.axhline(0, color="#E3E7ED", lw=0.8, zorder=0)
    ax.axvline(0, color="#E3E7ED", lw=0.8, zorder=0)
    return fig, ax


def save_panel(fig, name):
    path = tmp(name)
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)
    return path


def scalar_panel(f, size_in, levels):
    fig, ax = new_panel(size_in)
    xs = np.linspace(-LIM, LIM, 300)
    X, Y = np.meshgrid(xs, xs)
    Z = f(X, Y)
    ax.contourf(X, Y, Z, levels=40, cmap="Blues", alpha=0.85)
    ax.contour(X, Y, Z, levels=levels, colors=NAVY, linewidths=0.8, alpha=0.7,
               linestyles="solid")
    return save_panel(fig, "scalar")


def vector_panel(fx, fy, size_in, color, contour=None, levels=None):
    fig, ax = new_panel(size_in)
    if contour is not None:
        xs = np.linspace(-LIM, LIM, 300)
        X, Y = np.meshgrid(xs, xs)
        ax.contour(X, Y, contour(X, Y), levels=levels, colors="#9FB4CC", linewidths=0.8,
                   linestyles="solid")
    X, Y = np.meshgrid(GRID, GRID)
    U, V = fx(X, Y), fy(X, Y)
    # いちばん長い矢印がグリッド間隔の9割になるようにそろえる
    scale = np.hypot(U, V).max() / ((GRID[1] - GRID[0]) * 0.9)
    ax.quiver(X, Y, U, V, color=color, angles="xy", scale_units="xy", scale=scale,
              width=0.010, headwidth=4, headlength=4.5, headaxislength=4)
    return save_panel(fig, "vector")


def div_out_panel(value, size_in):
    fig, ax = new_panel(size_in)
    if value > 0:
        ax.add_patch(plt.Rectangle((-LIM, -LIM), 2 * LIM, 2 * LIM, color=LIGHT[TEAL], zorder=0))
        X, Y = np.meshgrid(GRID, GRID)
        ax.scatter(X, Y, marker="+", s=90, color=TEAL, linewidths=1.8)
    else:
        ax.add_patch(plt.Rectangle((-LIM, -LIM), 2 * LIM, 2 * LIM, color="#F3F4F6", zorder=0))
    return save_panel(fig, "div")


def rot_out_panel(nonzero, size_in):
    fig, ax = new_panel(size_in)
    if nonzero:
        X, Y = np.meshgrid(GRID, GRID)
        ax.scatter(X, Y, marker="o", s=150, facecolors="white", edgecolors=NAVY, linewidths=1.6)
        ax.scatter(X, Y, marker="o", s=14, color=NAVY)
    else:
        ax.add_patch(plt.Rectangle((-LIM, -LIM), 2 * LIM, 2 * LIM, color="#F3F4F6", zorder=0))
    return save_panel(fig, "rot")


# ---------- pptx ----------

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

ROW_Y = [0.35, 3.9]
ROW_H = 3.25
NODE_W = 5.3
IN_X = 0.45
OUT_X = 7.58
PANEL = 2.75


def add_text(slide, x, y, w, h, text, size, color, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = rgb(color)
    return tb


def add_math(slide, x, y, tex, size_pt=22, max_w=None, color=TEXT):
    path, w, h = mathimg(tex, size_pt, color)
    if max_w is not None and w > max_w:
        h = h * max_w / w
        w = max_w
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def node(slide, x, y, accent, kind, title, tex, panel_path, note):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                 Inches(NODE_W), Inches(ROW_H))
    box.adjustments[0] = 0.08
    box.fill.solid()
    box.fill.fore_color.rgb = rgb("#FFFFFF" if kind == "入力" else LIGHT[accent])
    box.line.color.rgb = rgb(accent)
    box.line.width = Pt(1.4)
    box.shadow.inherit = False
    slide.shapes.add_picture(panel_path, Inches(x + 0.15), Inches(y + 0.22),
                             Inches(PANEL), Inches(PANEL))
    tx = x + PANEL + 0.3
    tw = NODE_W - PANEL - 0.42
    add_text(slide, tx, y + 0.3, tw, 0.35, kind, 16, accent, bold=True)
    add_text(slide, tx, y + 0.65, tw, 0.35, title, 18, TEXT, bold=True)
    _, mh = add_math(slide, tx, y + 1.1, tex, 22, max_w=tw)
    add_text(slide, tx, y + 1.25 + mh, tw, 1.3, note, 15, GRAY)


def arrow(slide, y, accent, op_tex):
    x1 = IN_X + NODE_W + 0.12
    x2 = OUT_X - 0.12
    cy = y + ROW_H / 2
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(cy), Inches(x2), Inches(cy))
    c.line.color.rgb = rgb(accent)
    c.line.width = Pt(2.0)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "lg", "len": "lg"})
    ln.append(tail)
    path, w, h = mathimg(op_tex, 26, accent)
    slide.shapes.add_picture(path, Inches((x1 + x2) / 2 - w / 2), Inches(cy - h - 0.12),
                             Inches(w), Inches(h))


def row(slide, i, accent, op_tex, inp, out):
    y = ROW_Y[i]
    node(slide, IN_X, y, accent, "入力", *inp)
    node(slide, OUT_X, y, accent, "出力", *out)
    arrow(slide, y, accent, op_tex)


# 図1 勾配
s = prs.slides.add_slide(BLANK)
bowl = lambda X, Y: X**2 + Y**2
slope = lambda X, Y: X + Y
row(s, 0, NAVY, r"\nabla",
    ("スカラー場", r"f = x^2 + y^2", scalar_panel(bowl, PANEL, [0.5, 1, 2, 3, 4, 6]),
     "原点がいちばん低い\nお椀の形"),
    ("ベクトル場", r"\nabla f = (2x,\ 2y)",
     vector_panel(lambda X, Y: 2 * X, lambda X, Y: 2 * Y, PANEL, NAVY, bowl, [0.5, 1, 2, 3, 4, 6]),
     "坂を登る向きを指す\n外ほど急なので長い"))
row(s, 1, TEAL, r"\nabla",
    ("スカラー場", r"f = x + y", scalar_panel(slope, PANEL, np.arange(-4, 5, 1)),
     "右上ほど一様に高い\nまっすぐな坂"),
    ("ベクトル場", r"\nabla f = (1,\ 1)",
     vector_panel(lambda X, Y: np.ones_like(X), lambda X, Y: np.ones_like(Y), PANEL, TEAL,
                  slope, np.arange(-4, 5, 1)),
     "どこでも同じ向き\n同じ長さ"))

# 図2 発散
s = prs.slides.add_slide(BLANK)
row(s, 0, TEAL, r"\nabla \cdot",
    ("ベクトル場", r"F = (x,\ y,\ z)",
     vector_panel(lambda X, Y: X, lambda X, Y: Y, PANEL, TEAL),
     "原点から外へ広がる\nz = 0 の面で表示"),
    ("スカラー場", r"\nabla \cdot F = 3", div_out_panel(3, PANEL),
     "どこでも正\n＝湧き出し"))
row(s, 1, NAVY, r"\nabla \cdot",
    ("ベクトル場", r"F = (-y,\ x,\ 0)",
     vector_panel(lambda X, Y: -Y, lambda X, Y: X, PANEL, NAVY),
     "原点のまわりを\n反時計回り"),
    ("スカラー場", r"\nabla \cdot F = 0", div_out_panel(0, PANEL),
     "どこでも0\n＝湧き出しも\n　吸い込みもない"))

# 図3 回転
s = prs.slides.add_slide(BLANK)
row(s, 0, NAVY, r"\nabla \times",
    ("ベクトル場", r"F = (-y,\ x,\ 0)",
     vector_panel(lambda X, Y: -Y, lambda X, Y: X, PANEL, NAVY),
     "原点のまわりを\n反時計回り"),
    ("ベクトル場", r"\nabla \times F = (0,\ 0,\ 2)", rot_out_panel(True, PANEL),
     "⊙ は手前向き（+z）\n右ねじの向き"))
row(s, 1, TEAL, r"\nabla \times",
    ("ベクトル場", r"F = (x,\ y,\ z)",
     vector_panel(lambda X, Y: X, lambda X, Y: Y, PANEL, TEAL),
     "原点から外へ広がる\nz = 0 の面で表示"),
    ("ベクトル場", r"\nabla \times F = (0,\ 0,\ 0)", rot_out_panel(False, PANEL),
     "どこでも0\n＝渦がない"))

pptx_path = os.path.join(OUT, "ナブラ図.pptx")
prs.save(pptx_path)
print(pptx_path)
