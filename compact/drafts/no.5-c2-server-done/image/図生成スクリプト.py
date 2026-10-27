import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

NAVY=RGBColor(0x2E,0x5E,0x8E); TEAL=RGBColor(0x1E,0x8A,0x82); RED=RGBColor(0xC0,0x39,0x2B)
TXT=RGBColor(0x1F,0x3A,0x5C); GRAY=RGBColor(0x6B,0x72,0x80); SLATE=RGBColor(0x2B,0x3A,0x4A)
WHITE=RGBColor(0xFF,0xFF,0xFF); LNAVY=RGBColor(0xEE,0xF3,0xF9); LTEAL=RGBColor(0xE8,0xF5,0xF3)
LRED=RGBColor(0xFB,0xEE,0xEC); LGRAY=RGBColor(0xF3,0xF4,0xF6); MGRAY=RGBColor(0x9A,0xA5,0xB4)
FONT="Noto Sans CJK JP"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6])
bg=s.background.fill; bg.solid(); bg.fore_color.rgb=WHITE

def text(shape, lines, size=14, color=TXT, bold=False, align=PP_ALIGN.CENTER):
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(0.05); tf.margin_top=tf.margin_bottom=Inches(0.02)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    if isinstance(lines,str): lines=[(lines,size,color,bold)]
    for i,(t,sz,c,b) in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=t; f=r.font; f.size=Pt(sz); f.color.rgb=c; f.bold=b; f.name=FONT
        rPr=r._r.get_or_add_rPr(); ea=rPr.find(qn('a:ea'))
        if ea is None: ea=etree.SubElement(rPr,qn('a:ea'))
        ea.set('typeface',FONT)

def box(x,y,w,h,fill=WHITE,line=NAVY,lw=1.4,shape=MSO_SHAPE.ROUNDED_RECTANGLE,adj=0.16):
    sh=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=adj
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False
    return sh

def label(x,y,w,h,lines,size=12,color=GRAY,bold=False,align=PP_ALIGN.CENTER):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); text(tb,lines,size,color,bold,align); return tb

def line(x1,y1,x2,y2,color=MGRAY,w=1.5,arrow=True,dash=False):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w)
    ln=c.line._get_or_add_ln()
    if dash:
        d=etree.SubElement(ln,qn('a:prstDash')); d.set('val','dash')
    if arrow:
        t=etree.SubElement(ln,qn('a:tailEnd')); t.set('type','triangle'); t.set('w','lg'); t.set('len','lg')
    return c

def zone(x,y,w,h,tab,accent,fill,size=18):
    box(x,y,w,h,fill=fill,line=accent,lw=1.8,adj=0.05)
    tw=size/72*1.05*len(tab)+0.6
    t=box(x+0.25,y-0.27,tw,0.54,fill=accent,line=None,adj=0.5); text(t,tab,size,WHITE,True)

def node(x,y,w,h,title,sub,accent,fill=WHITE,lw=2.0):
    b=box(x,y,w,h,fill=fill,line=accent,lw=lw)
    text(b,[(title,22,TXT,True)]+[(t,16,GRAY,False) for t in sub])
    return b

def pill(cx,cy,t,fill,size=16,color=WHITE,w=None):
    w=w or size/72*1.05*len(t)+0.5
    p=box(cx-w/2,cy-0.25,w,0.5,fill=fill,line=None,adj=0.5); text(p,t,size,color,True); return p

ZT,ZB=0.95,6.75; ZH=ZB-ZT
# zones
zone(0.3,ZT,6.05,ZH,"インターネット",SLATE,LGRAY)
zone(6.85,ZT,6.15,ZH,"社内ネットワーク",NAVY,LNAVY)

NT,NH=1.85,2.45
node(0.6,NT,1.8,NH,"攻撃者",["操作する人"],RED,fill=LRED)
node(3.5,NT,2.4,NH,"C2サーバ",["命令の置き場","データの受け皿"],RED,lw=2.4)
node(10.2,NT,2.5,NH,"感染端末",["マルウェアが","常駐"],RED,lw=2.4)
node(10.35,5.35,2.2,1.0,"ほかの端末",[],NAVY)

# attacker <-> C2
line(2.4,2.75,3.5,2.75,color=RED,w=1.8); label(2.4,2.27,1.1,0.45,"命令",16,RED,True)
line(3.5,3.65,2.4,3.65,color=RED,w=1.8); label(2.4,3.17,1.1,0.45,"結果",16,RED,True)

# firewall wall
WX=6.6
box(WX-0.17,NT-0.1,0.34,3.9,fill=SLATE,line=None,adj=0.5)
pill(WX,1.4,"FW・プロキシ",SLATE,size=16)

# lanes C2 <-> infected PC
L=[(2.3,"① ビーコン「命令ある？」",False),(3.1,"② 命令・追加のマルウェア",True),(3.9,"③ 実行結果・盗んだデータ",False)]
for y,t,toPC in L:
    if toPC: line(5.9,y,10.2,y,color=RED,w=2.0)
    else: line(10.2,y,5.9,y,color=RED,w=2.0)
    label(6.85,y-0.45,3.3,0.42,t,16,RED,True)
label(10.2,1.3,2.5,0.45,"①〜③をくり返す",16,GRAY,True)

# blocked inbound
line(4.7,NT+NH,4.7,5.0,color=MGRAY,w=1.6,arrow=False,dash=True)
line(4.7,5.0,WX-0.2,5.0,color=MGRAY,w=1.6,dash=True)
label(WX-0.55,4.7,0.6,0.6,"✘",26,RED,True)
label(2.9,5.1,3.5,0.45,"外→内の接続は止められる",16,GRAY,True)

# lateral movement
line(11.45,NT+NH,11.45,5.35,color=RED,w=1.8,dash=True)
label(11.55,4.6,1.4,0.42,"④ 横展開",16,RED,True,PP_ALIGN.LEFT)

# trace tags
pill(8.6,5.0,"痕跡：EDRログ",TEAL,size=16)
line(9.4,4.75,10.2,4.15,color=TEAL,w=1.5)
pill(WX,6.25,"痕跡：通信・DNSログ",TEAL,size=16)
line(WX,6.0,WX,NT+3.8+0.05,color=TEAL,w=1.5)

prs.save(sys.argv[1])
