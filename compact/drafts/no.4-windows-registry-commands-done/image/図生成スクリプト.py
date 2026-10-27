import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

NAVY=RGBColor(0x2E,0x5E,0x8E); TEAL=RGBColor(0x1E,0x8A,0x82); RED=RGBColor(0xC0,0x39,0x2B)
TXT=RGBColor(0x1F,0x3A,0x5C); GRAY=RGBColor(0x6B,0x72,0x80); SLATE=RGBColor(0x2B,0x3A,0x4A)
WHITE=RGBColor(0xFF,0xFF,0xFF); LNAVY=RGBColor(0xEE,0xF3,0xF9); LTEAL=RGBColor(0xE8,0xF5,0xF3)
LSLATE=RGBColor(0xF1,0xF3,0xF6); LGRAY=RGBColor(0x9A,0xA5,0xB4)
FONT="Noto Sans CJK JP"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(11.3)
s=prs.slides.add_slide(prs.slide_layouts[6])
bg=s.background.fill; bg.solid(); bg.fore_color.rgb=WHITE

def text(shape, lines, size=14, color=TXT, bold=False, align=PP_ALIGN.CENTER):
    tf=shape.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(0.08); tf.margin_top=tf.margin_bottom=Inches(0.03)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    if isinstance(lines,str): lines=[(lines,size,color,bold)]
    lines=[(l,size,color,bold) if isinstance(l,str) else l for l in lines]
    for i,(t,sz,c,b) in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=t; f=r.font; f.size=Pt(sz); f.color.rgb=c; f.bold=b; f.name=FONT
        rPr=r._r.get_or_add_rPr(); ea=rPr.find(qn('a:ea'))
        if ea is None: ea=etree.SubElement(rPr,qn('a:ea'))
        ea.set('typeface',FONT)

def box(x,y,w,h,fill=WHITE,line=NAVY,lw=1.4,shape=MSO_SHAPE.ROUNDED_RECTANGLE,adj=0.16,dash=False):
    sh=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=adj
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb=line; sh.line.width=Pt(lw)
        if dash:
            d=etree.SubElement(sh.line._get_or_add_ln(),qn('a:prstDash')); d.set('val','dash')
    sh.shadow.inherit=False
    return sh

def label(x,y,w,h,lines,size=13,color=GRAY,bold=False,align=PP_ALIGN.CENTER):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); text(tb,lines,size,color,bold,align); return tb

def line(x1,y1,x2,y2,color=NAVY,w=1.5,arrow=True,dash=False,both=False):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w)
    ln=c.line._get_or_add_ln()
    if dash:
        d=etree.SubElement(ln,qn('a:prstDash')); d.set('val','dash')
    if both:
        h=etree.SubElement(ln,qn('a:headEnd')); h.set('type','triangle'); h.set('w','lg'); h.set('len','lg')
    if arrow:
        t=etree.SubElement(ln,qn('a:tailEnd')); t.set('type','triangle'); t.set('w','lg'); t.set('len','lg')
    return c

def zone(x,y,w,h,tab,accent,fill,size=18):
    box(x,y,w,h,fill=fill,line=accent,lw=1.8,adj=0.04)
    tw=size/72*1.05*len(tab)+0.6
    t=box(x+0.25,y-0.27,tw,0.54,fill=accent,line=None,adj=0.5); text(t,tab,size,WHITE,True)

def node(x,y,w,h,title,subs=(),accent=NAVY,fill=WHITE,tsize=16,ssize=13,dash=False):
    b=box(x,y,w,h,fill=fill,line=accent,lw=1.4,dash=dash)
    text(b,[(title,tsize,TXT,True)]+[(t,ssize,GRAY,False) for t in subs]); return b

# ---------------- ユーザーモード ----------------
zone(0.35,0.45,12.65,3.95,"ユーザーモード",NAVY,LNAVY)
apps=[("regedit.exe","GUI で見て編集する"),("reg.exe","コマンドで読み書き"),("PowerShell","スクリプトで操作"),("アプリ","自分の設定を保存")]
AX=[4.35,6.45,8.55,10.65]; AW=1.95
for x,(t,sub) in zip(AX,apps):
    node(x,0.85,AW,0.9,t,[sub])
    line(x+AW/2,1.75,x+AW/2,2.1)
node(4.35,2.1,8.25,0.95,"レジストリ API（advapi32.dll / KernelBase.dll）",
     ["アプリが呼ぶ窓口。キーを開く・値を読む・書く関数を提供（RegOpenKeyEx など）","HKCU などのルートキーを、カーネルが分かるパスに変換する"])
line(8.475,3.05,8.475,3.35)
node(4.35,3.35,8.25,0.75,"ntdll.dll",["API の呼び出しをシステムコールにして、カーネルへ渡す（NtOpenKey など）"])

# ルートキーの変換表
mb=box(0.6,0.85,3.45,3.25,fill=WHITE,line=NAVY,lw=1.4,adj=0.06)
L=PP_ALIGN.LEFT
tf_lines=[("ルートキーの変換先",16,TXT,True),
 ("HKLM → \\Registry\\Machine",13,TXT,False),
 ("HKU → \\Registry\\User",13,TXT,False),
 ("HKCU → \\Registry\\User\\<SID>",13,TXT,False),
 ("HKCR → 次の2つを合成",13,TXT,False),
 ("　HKLM\\SOFTWARE\\Classes",13,GRAY,False),
 ("　HKU\\<SID>_Classes",13,GRAY,False),
 ("HKCC → HKLM\\SYSTEM 内の",13,TXT,False),
 ("　現在のハードウェアプロファイル",13,GRAY,False)]
text(mb,tf_lines,align=L)
mb.text_frame.margin_left=Inches(0.2)
line(4.35,2.575,4.05,2.575,dash=True)

# ---------------- カーネルモード ----------------
zone(0.35,4.85,12.65,3.2,"カーネルモード",SLATE,LSLATE)
line(6.15,4.1,6.15,5.3)
label(6.25,4.5,2.2,0.35,"システムコール",14,SLATE,True,L)
node(0.6,5.3,2.9,1.05,"カーネルドライバ",["自分の設定などを","カーネル内から直接読み書き"],accent=SLATE)
node(4.35,5.3,3.4,1.05,"オブジェクトマネージャー",["パスをたどって目的のキーを探し、","開いたキーにハンドルを発行する"],accent=SLATE)
node(9.25,5.3,3.35,1.05,"構成マネージャー",["ntoskrnl.exe 内のレジストリ本体","キーを探し、値を読み書きする"],accent=SLATE)
line(3.5,5.825,4.35,5.825,color=SLATE)
label(3.45,5.3,0.95,0.42,"Zw* 関数",13,GRAY)
line(7.75,5.825,9.25,5.825,color=SLATE)
label(7.75,5.05,1.5,0.7,[("名前解決",13,GRAY,False),("ACL チェック",13,GRAY,False)])
# 下段
node(0.6,6.75,7.35,1.1,"レジストリ コールバック（CmRegisterCallbackEx）",
     ["キーの作成や値の変更を、登録したドライバ（EDR・Sysmon など）に知らせる"],accent=SLATE,dash=True)
line(9.25,6.25,7.95,7.0,color=SLATE,dash=True)
label(8.0,6.25,1.0,0.35,"通知",13,GRAY)
line(10.925,6.35,10.925,6.75,color=SLATE)
node(9.25,6.75,3.35,1.1,"メモリ上のハイブ",
     ["キーと値をメモリに載せて","高速に読み書きする","変更はあとでディスクへ書き戻す"],accent=SLATE)

# ---------------- ディスク ----------------
zone(0.35,8.55,12.65,2.55,"ディスク（ハイブファイル）",TEAL,LTEAL)
line(10.925,7.85,10.925,8.55,color=TEAL,both=True)
label(5.6,8.08,5.25,0.4,"読み込み・書き戻し（I/O マネージャー経由）",13,GRAY,False,PP_ALIGN.RIGHT)
label(0.6,8.95,8.1,0.38,"%SystemRoot%\\System32\\config\\",14,TEAL,True)
label(8.85,8.95,1.95,0.38,"%UserProfile%\\",14,TEAL,True)
label(10.9,8.78,2.05,0.6,["%LocalAppData%\\","Microsoft\\Windows\\"],13,TEAL,True)
files=[(0.6,1.95,"SYSTEM",["起動・ドライバ・","サービスの設定"],["→ HKLM\\SYSTEM"]),
       (2.65,1.95,"SOFTWARE",["マシン全体の","ソフト・OS の設定"],["→ HKLM\\SOFTWARE"]),
       (4.7,1.95,"SAM, SECURITY",["アカウント・認証情報"],["→ HKLM\\SAM","→ HKLM\\SECURITY"]),
       (6.75,1.95,"DEFAULT",["SYSTEM 用の","ユーザー設定"],["→ HKU\\.DEFAULT"]),
       (8.85,1.95,"NTUSER.DAT",["HKCU の実体","ユーザーごとの設定"],["→ HKU\\<SID>"]),
       (10.9,2.05,"UsrClass.dat",["ユーザーごとの","関連付け・COM"],["→ HKU\\<SID>_Classes"])]
for x,w,t,role,keys in files:
    b=box(x,9.35,w,1.3,fill=WHITE,line=TEAL,lw=1.4)
    text(b,[(t,16,TXT,True)]+[(r,12,GRAY,False) for r in role]+[(k,12,TEAL,False) for k in keys])
label(0.6,10.68,12.35,0.35,"各ハイブには .LOG1 / .LOG2 のトランザクションログが付き、書き込み途中で落ちても整合性を保てる",13,GRAY)

prs.save(sys.argv[1])
