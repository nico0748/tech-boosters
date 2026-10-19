import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

OUT = sys.argv[1]
FONT = "Hiragino Sans"
MONO = "Menlo"
NAVY = RGBColor(0x24, 0x40, 0x4D)
SUB = RGBColor(0x5E, 0x76, 0x82)
GRAY_BG = RGBColor(0xF1, 0xF3, 0xF4)
GRAY_LN = RGBColor(0xC2, 0xCB, 0xD0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# LangChain = ティール、LangGraph = 青、LangSmith = オレンジ、Langfuse = 赤系、AWS = グレー
LC = RGBColor(0x1A, 0xA6, 0x9A)
LC_D = RGBColor(0x0E, 0x7C, 0x72)
LC_BG = RGBColor(0xE6, 0xF5, 0xF2)
LG = RGBColor(0x2F, 0x6F, 0xB5)
LG_D = RGBColor(0x1F, 0x4F, 0x86)
LG_BG = RGBColor(0xE8, 0xF0, 0xFA)
LS = RGBColor(0xE0, 0x7B, 0x1F)
LS_D = RGBColor(0xA8, 0x55, 0x0C)
LS_BG = RGBColor(0xFD, 0xF1, 0xE5)
LF = RGBColor(0xC8, 0x46, 0x5A)
LF_D = RGBColor(0x9A, 0x2E, 0x40)
LF_BG = RGBColor(0xFB, 0xEB, 0xEE)

IN = 914400
prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]


def inch(v):
    return Emu(int(v * IN))


def _runs(para, segs, size, color, bold, font):
    for seg, opt in segs:
        r = para.add_run()
        r.text = seg
        f = opt.get("font", font)
        r.font.name = f
        r.font.size = Pt(opt.get("size", size))
        r.font.bold = opt.get("bold", bold)
        r.font.color.rgb = opt.get("color", color)
        rpr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rpr, qn(tag))
            el.set("typeface", f)


def fill_tf(tf, lines, size, color, bold, align, font, anchor):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(int(0.06 * IN))
    tf.margin_top = tf.margin_bottom = Emu(int(0.03 * IN))
    lines = lines.split("\n") if isinstance(lines, str) else lines
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = line if isinstance(line, list) else [(line, {})]
        _runs(para, segs, size, color, bold, font)


def text(slide, x, y, w, h, s, size=14, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         font=FONT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    fill_tf(tb.text_frame, s, size, color, bold, align, font, anchor)
    return tb


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.5,
        radius=0.08, dash=False):
    sh = slide.shapes.add_shape(shape, inch(x), inch(y), inch(w), inch(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            ln = sh.line._get_or_add_ln()
            pd = etree.SubElement(ln, qn("a:prstDash"))
            pd.set("val", "dash")
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def node(slide, x, y, w, h, title, sub=None, fill=WHITE, line=NAVY, tc=NAVY, size=15,
         sub_size=12, shape=MSO_SHAPE.ROUNDED_RECTANGLE, title_font=FONT, lw=1.75):
    sh = box(slide, x, y, w, h, fill, line, shape=shape, lw=lw)
    lines = [[(title, {"bold": True, "font": title_font})]]
    if sub:
        for sl in sub.split("\n"):
            lines.append([(sl, {"size": sub_size, "color": SUB})])
    fill_tf(sh.text_frame, lines, size, tc, False, PP_ALIGN.CENTER, FONT, MSO_ANCHOR.MIDDLE)
    return sh


def frame(slide, x, y, w, h, label, color, bg, label_w=None, dash=False):
    box(slide, x, y, w, h, bg, color, lw=2.0, radius=0.03, dash=dash)
    lw = label_w or (len(label) * 0.2 + 0.3)
    tag = box(slide, x + 0.18, y - 0.2, lw, 0.4, color, radius=0.25)
    fill_tf(tag.text_frame, [[(label, {"bold": True})]], 14, WHITE, True, PP_ALIGN.CENTER,
            FONT, MSO_ANCHOR.MIDDLE)
    return tag


def arrow(slide, x1, y1, x2, y2, color=SUB, w=2.0, both=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", "dash")
    if both:
        head = etree.SubElement(ln, qn("a:headEnd"))
        head.set("type", "triangle")
        head.set("w", "med")
        head.set("len", "med")
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def alabel(slide, cx, y, s, w=None, color=SUB, size=12, align=PP_ALIGN.CENTER):
    """矢印の注記。y は矢印の y 座標。注記は矢印の上に置く"""
    w = w or (len(s) * 0.19 + 0.2)
    return text(slide, cx - w / 2, y - 0.36, w, 0.3, s, size, color, align=align,
                anchor=MSO_ANCHOR.BOTTOM)


def side_label(slide, x, cy, s, w=None, color=SUB, size=12):
    """縦向きの矢印の注記（矢印の右側）"""
    w = w or (len(s) * 0.19 + 0.2)
    return text(slide, x, cy - 0.15, w, 0.3, s, size, color, anchor=MSO_ANCHOR.MIDDLE)


def chip(slide, x, y, w, h, s, fill, line, tc, size=13, font=FONT, bold=True):
    sh = box(slide, x, y, w, h, fill, line, lw=1.25, radius=0.2)
    fill_tf(sh.text_frame, [[(s, {"bold": bold, "font": font})]], size, tc, bold,
            PP_ALIGN.CENTER, font, MSO_ANCHOR.MIDDLE)
    return sh


def oval(slide, x, y, w, h, s, fill=NAVY, tc=WHITE, size=13):
    sh = box(slide, x, y, w, h, fill, None, shape=MSO_SHAPE.OVAL)
    fill_tf(sh.text_frame, [[(s, {"bold": True, "font": MONO})]], size, tc, True,
            PP_ALIGN.CENTER, MONO, MSO_ANCHOR.MIDDLE)
    return sh


# =====================================================================
# 図1：LangChain 単体
# =====================================================================
s = prs.slides.add_slide(BLANK)

node(s, 0.4, 1.3, 1.5, 4.6, "ユーザー", "質問・指示", fill=GRAY_BG, line=GRAY_LN)

frame(s, 2.7, 0.75, 6.3, 5.95, "LangChain（create_agent）", LC, LC_BG, label_w=3.4)
node(s, 3.1, 1.45, 2.3, 1.15, "プロンプト", "テンプレートに\n入力を埋め込む", line=LC, tc=LC_D)
node(s, 6.3, 1.45, 2.3, 1.15, "チャットモデル", "共通の書き方で\nモデルを呼ぶ", line=LC, tc=LC_D)
node(s, 6.3, 4.85, 2.3, 1.15, "ツール", "@tool で定義した\n関数", line=LC, tc=LC_D)
node(s, 3.1, 4.85, 2.3, 1.15, "出力", "テキスト・\n構造化出力", line=LC, tc=LC_D)

# ユーザー → プロンプト → モデル
arrow(s, 1.9, 2.03, 3.1, 2.03)
alabel(s, 2.33, 2.03, "入力", w=0.6)
arrow(s, 5.4, 2.03, 6.3, 2.03)
# モデル ⇄ ツール（ループ）
arrow(s, 7.15, 2.6, 7.15, 4.85, color=LC_D)
arrow(s, 7.75, 4.85, 7.75, 2.6, color=LC_D)
side_label(s, 5.75, 3.72, "ツール呼び出し", w=1.35, color=LC_D)
side_label(s, 7.85, 3.72, "結果", w=0.7, color=LC_D)
# モデル → 出力 → ユーザー
arrow(s, 6.3, 2.45, 5.0, 4.85)
arrow(s, 3.1, 5.43, 1.9, 5.43)
alabel(s, 2.33, 5.43, "回答", w=0.6)

# 外部：LLMプロバイダ
frame(s, 9.85, 0.75, 3.05, 2.85, "LLMプロバイダ", GRAY_LN, GRAY_BG, label_w=1.9)
for i, name in enumerate(["OpenAI", "Anthropic", "Google", "ほか"]):
    chip(s, 10.2, 1.2 + i * 0.57, 2.35, 0.45, name, WHITE, GRAY_LN, NAVY)
arrow(s, 8.6, 1.8, 9.85, 1.8, both=True)
alabel(s, 9.22, 1.8, "API")

# 外部：ツールの先
frame(s, 9.85, 4.15, 3.05, 2.55, "外部のサービス", GRAY_LN, GRAY_BG, label_w=2.0)
for i, name in enumerate(["Web検索", "社内API", "データベース"]):
    chip(s, 10.2, 4.6 + i * 0.6, 2.35, 0.47, name, WHITE, GRAY_LN, NAVY)
arrow(s, 8.6, 5.43, 9.85, 5.43, both=True)
alabel(s, 9.22, 5.43, "実行")

# =====================================================================
# 図2：LangChain + LangGraph のオーケストレーション
# =====================================================================
s = prs.slides.add_slide(BLANK)

frame(s, 0.4, 0.55, 9.3, 6.6, "LangGraph（StateGraph）", LG, LG_BG, label_w=3.3)
oval(s, 0.6, 2.98, 1.2, 0.55, "START", size=12)
node(s, 2.2, 2.6, 2.3, 1.3, "スーパーバイザー", "次の担当を決める", line=LG, tc=LG_D)
arrow(s, 1.8, 3.25, 2.2, 3.25, color=LG_D)


def agent_node(y, title):
    sh = box(s, 5.35, y, 2.65, 1.25, WHITE, LG, lw=1.75)
    fill_tf(sh.text_frame, [[(title, {"bold": True})]], 15, LG_D, True, PP_ALIGN.CENTER, FONT,
            MSO_ANCHOR.TOP)
    sh.text_frame.margin_top = Emu(int(0.12 * IN))
    chip(s, 5.55, y + 0.62, 2.25, 0.45, "LangChainで実装", LC_BG, LC, LC_D, size=12)


agent_node(1.0, "調査エージェント")
agent_node(2.65, "執筆エージェント")
node(s, 5.35, 4.55, 2.65, 1.15, "人の確認", "interrupt で一時停止", line=LG, tc=LG_D)
oval(s, 6.18, 6.25, 1.0, 0.55, "END")

# スーパーバイザー ⇄ エージェント（条件分岐）
arrow(s, 4.5, 2.85, 5.35, 1.62, color=LG_D, both=True)
arrow(s, 4.5, 3.25, 5.35, 3.25, color=LG_D, both=True)
arrow(s, 4.5, 3.65, 5.35, 5.05, color=LG_D)
text(s, 2.2, 4.05, 2.3, 0.6, "条件付きエッジで\n行き先を分岐", 12, SUB, align=PP_ALIGN.CENTER)
arrow(s, 6.68, 5.7, 6.68, 6.25, color=LG_D)
side_label(s, 6.78, 5.97, "承認されたら", w=1.3)

# State
sb = box(s, 8.25, 1.0, 1.2, 3.6, WHITE, LG, lw=1.5, dash=True)
fill_tf(sb.text_frame, [[("State", {"bold": True, "font": MONO})], [("共有の", {"size": 12})],
                        [("状態", {"size": 12})], [("", {"size": 12})],
                        [("会話履歴", {"size": 12, "color": SUB})],
                        [("途中の結果", {"size": 12, "color": SUB})]],
        15, LG_D, True, PP_ALIGN.CENTER, FONT, MSO_ANCHOR.MIDDLE)

# チェックポイント保存先
node(s, 10.45, 1.55, 2.5, 1.5, "チェックポインタ", "保存先：Postgres・\nSQLite・メモリなど",
     fill=GRAY_BG, line=GRAY_LN)
arrow(s, 9.7, 2.3, 10.45, 2.3, both=True)
alabel(s, 10.07, 2.3, "保存", w=0.7)
text(s, 10.45, 3.15, 2.5, 0.6, "失敗しても途中から再開", 12, SUB, align=PP_ALIGN.CENTER)

# 人
node(s, 10.45, 4.55, 2.5, 1.15, "人（承認する人）", "内容を確認", fill=GRAY_BG, line=GRAY_LN)
arrow(s, 10.45, 5.12, 8.0, 5.12)
alabel(s, 9.1, 5.12, "承認・修正", w=1.2)

# =====================================================================
# 図3：RAG（LangGraph の agentic RAG）
# =====================================================================
s = prs.slides.add_slide(BLANK)

# 事前準備（インデックス作成）
frame(s, 0.4, 0.55, 7.25, 1.75, "事前準備（LangChainの部品）", LC, LC_BG, label_w=3.5)
prep = [("文書の読み込み", "Web・PDFなど"), ("テキスト分割", "小さく分ける"),
        ("埋め込みモデル", "ベクトルに変換")]
px = [0.65, 2.45, 4.25]
for (t, sub), x in zip(prep, px):
    node(s, x, 0.95, 1.5, 1.1, t, sub, line=LC, tc=LC_D, size=13)
for x in px:
    arrow(s, x + 1.5, 1.5, x + 1.8, 1.5, color=LC_D)
node(s, 6.05, 0.95, 1.4, 1.1, "ベクトル", "ストア", line=LC, tc=LC_D, size=14, sub_size=14)

# 質問が来たとき（LangGraph）
frame(s, 0.4, 2.9, 12.5, 4.25, "LangGraph（質問が来たとき）", LG, LG_BG, label_w=3.6)
oval(s, 0.55, 4.47, 1.2, 0.55, "START", size=12)
node(s, 2.05, 4.1, 2.35, 1.3, "検索するか判断", "モデル＋検索ツール", line=LG, tc=LG_D)
node(s, 5.05, 4.1, 2.0, 1.3, "検索", "ToolNode", line=LG, tc=LG_D)
d = box(s, 7.75, 3.95, 2.2, 1.6, WHITE, LG, shape=MSO_SHAPE.DIAMOND, lw=1.75)
fill_tf(d.text_frame, [[("関連ある？", {"bold": True})]], 14, LG_D, True, PP_ALIGN.CENTER,
        FONT, MSO_ANCHOR.MIDDLE)
node(s, 10.65, 4.1, 2.0, 1.3, "回答を生成", "文書をもとに答える", line=LG, tc=LG_D)
node(s, 7.85, 6.05, 2.0, 0.9, "質問を書き換え", None, line=LG, tc=LG_D, size=14)
oval(s, 2.72, 3.25, 1.0, 0.5, "END")
oval(s, 11.15, 6.2, 1.0, 0.55, "END")

arrow(s, 1.75, 4.75, 2.05, 4.75, color=LG_D)
arrow(s, 4.4, 4.75, 5.05, 4.75, color=LG_D)
alabel(s, 4.72, 4.75, "検索", w=0.65)
arrow(s, 7.05, 4.75, 7.75, 4.75, color=LG_D)
arrow(s, 9.95, 4.75, 10.65, 4.75, color=LG_D)
alabel(s, 10.3, 4.75, "ある", w=0.65)
arrow(s, 8.85, 5.55, 8.85, 6.05, color=LG_D)
side_label(s, 8.95, 5.8, "ない", w=0.6)
arrow(s, 7.85, 6.5, 3.6, 5.4, color=LG_D)
text(s, 4.6, 6.15, 2.3, 0.3, "もう一度検索", 12, SUB, align=PP_ALIGN.CENTER)
arrow(s, 3.22, 4.1, 3.22, 3.75, color=LG_D)
side_label(s, 3.32, 3.92, "検索不要", w=1.0)
arrow(s, 11.65, 5.4, 11.65, 6.2, color=LG_D)
# 検索 ⇄ ベクトルストア
arrow(s, 6.05, 4.1, 6.75, 2.05, color=LC_D, both=True, dash=True)
side_label(s, 6.75, 2.65, "類似検索", w=1.0, color=LC_D)

# =====================================================================
# 図4：LangChain + LangGraph + LangSmith
# =====================================================================
s = prs.slides.add_slide(BLANK)

frame(s, 0.4, 0.9, 4.6, 5.7, "あなたのアプリ", NAVY, GRAY_BG, label_w=2.1)
box(s, 0.7, 1.4, 4.0, 3.15, LG_BG, LG, lw=1.5)
text(s, 0.85, 1.5, 3.7, 0.35, "LangGraph", 15, LG_D, bold=True, font=MONO)
oval(s, 0.95, 2.05, 0.6, 0.6, "", fill=LG)
oval(s, 2.4, 2.05, 0.6, 0.6, "", fill=LG)
oval(s, 3.85, 2.05, 0.6, 0.6, "", fill=LG)
arrow(s, 1.55, 2.35, 2.4, 2.35, color=LG_D)
arrow(s, 3.0, 2.35, 3.85, 2.35, color=LG_D)
text(s, 0.85, 2.75, 3.7, 0.3, "ノード・エッジ・State", 12, SUB, align=PP_ALIGN.CENTER)
chip(s, 0.95, 3.35, 3.5, 0.75, "LangChain（モデル・ツール）", LC_BG, LC, LC_D, size=13)
chip(s, 0.7, 4.95, 4.0, 1.1, "LangSmith SDK", LS_BG, LS, LS_D, size=15)
text(s, 0.7, 6.1, 4.0, 0.35, "LANGSMITH_TRACING=true", 12, SUB, font=MONO, align=PP_ALIGN.CENTER)

frame(s, 6.3, 0.9, 6.6, 4.25, "LangSmith（LangChain社のクラウド）", LS, LS_BG, label_w=4.3)
feats = [("トレース", "入出力・ツール呼び出し\n時間・トークン・費用"),
         ("評価", "データセット\nLLM-as-a-judge"),
         ("監視", "ダッシュボード\nアラート"),
         ("プロンプト管理", "バージョン管理\nPlayground")]
for i, (t, sub) in enumerate(feats):
    x = 6.65 + (i % 2) * 3.1
    y = 1.45 + (i // 2) * 1.75
    node(s, x, y, 2.85, 1.5, t, sub, line=LS, tc=LS_D)

arrow(s, 5.0, 2.3, 6.3, 2.3, color=LS_D)
alabel(s, 5.65, 2.3, "トレース送信", w=1.3, color=LS_D)
arrow(s, 6.3, 4.4, 5.0, 4.4, color=LS_D, dash=True)
alabel(s, 5.65, 4.4, "プロンプト取得", w=1.4, color=LS_D)

node(s, 8.3, 5.75, 2.6, 1.0, "開発者", "ブラウザで確認", fill=GRAY_BG, line=GRAY_LN)
arrow(s, 9.6, 5.75, 9.6, 5.15, both=True)

# =====================================================================
# 図5：LangChain + LangGraph + Langfuse（AWS セルフホスト / クラウド版）
# =====================================================================
s = prs.slides.add_slide(BLANK)

frame(s, 0.3, 0.9, 2.25, 5.9, "あなたのアプリ", NAVY, GRAY_BG, label_w=2.1)
chip(s, 0.5, 1.5, 1.85, 0.75, "LangGraph", LG_BG, LG, LG_D, size=14, font=MONO)
chip(s, 0.5, 2.45, 1.85, 0.75, "LangChain", LC_BG, LC, LC_D, size=14, font=MONO)
sdk = box(s, 0.5, 3.55, 1.85, 1.35, LF_BG, LF, lw=1.25, radius=0.15)
fill_tf(sdk.text_frame, [[("Langfuse SDK", {"bold": True})],
                         [("CallbackHandler", {"size": 12, "font": MONO, "color": LF_D})]],
        14, LF_D, True, PP_ALIGN.CENTER, FONT, MSO_ANCHOR.MIDDLE)
text(s, 0.35, 5.05, 2.15, 0.9, "送信先は\nLANGFUSE_BASE_URL\nで切り替え", 12, SUB, align=PP_ALIGN.CENTER)

# --- AWS にセルフホスト
frame(s, 3.75, 0.55, 9.2, 4.1, "AWS（自分のアカウントでセルフホスト）", GRAY_LN, WHITE, label_w=4.6)
row_y, row_h = 1.15, 1.25
aws = [(3.95, 1.3, "ALB", "入口", GRAY_BG, GRAY_LN, NAVY),
       (5.6, 1.6, "Langfuse Web", "UI・API\n（EKS上）", LF_BG, LF, LF_D),
       (7.55, 1.45, "Amazon S3", "受け取った\nイベント", GRAY_BG, GRAY_LN, NAVY),
       (9.35, 1.6, "Langfuse Worker", "非同期で処理\n（EKS上）", LF_BG, LF, LF_D),
       (11.3, 1.5, "ClickHouse", "トレース・\nスコア", GRAY_BG, GRAY_LN, NAVY)]
for x, w, t, sub, bg, ln, tc in aws:
    node(s, x, row_y, w, row_h, t, sub, fill=bg, line=ln, tc=tc, size=14)
cy = row_y + row_h / 2
for i in range(4):
    x, w = aws[i][0], aws[i][1]
    arrow(s, x + w, cy, aws[i + 1][0], cy)
node(s, 5.6, 3.15, 1.6, 1.15, "Aurora", "PostgreSQL\nユーザー・設定", fill=GRAY_BG, line=GRAY_LN,
     size=14)
node(s, 7.55, 3.15, 1.45, 1.15, "ElastiCache", "Redis\nキュー", fill=GRAY_BG, line=GRAY_LN,
     size=14)
arrow(s, 6.4, row_y + row_h, 6.4, 3.15)
arrow(s, 6.9, row_y + row_h, 7.75, 3.15)
arrow(s, 9.0, 3.4, 9.6, row_y + row_h)
text(s, 9.5, 3.2, 3.3, 0.6, "ClickHouse は EKS 上で動かす\n（ClickHouse Cloud も可）", 12, SUB)

# --- クラウド版
frame(s, 3.75, 5.2, 9.2, 1.85, "Langfuse Cloud（Web版）", LF, LF_BG, label_w=3.4)
regions = ["EU（アイルランド）", "US（オレゴン）", "日本（東京）", "HIPAA（米国）"]
for i, r in enumerate(regions):
    chip(s, 4.05 + i * 2.2, 5.65, 2.0, 0.6, r, WHITE, LF, LF_D, size=13)
text(s, 4.05, 6.4, 8.6, 0.35, "サーバーの運用は Langfuse 側。例：jp.cloud.langfuse.com", 12, SUB,
     align=PP_ALIGN.CENTER)

# アプリ → 送信先
arrow(s, 2.55, cy, 3.95, cy, color=LF_D)
alabel(s, 3.15, cy, "自前で運用", w=1.0, color=LF_D)
arrow(s, 2.55, 5.95, 3.75, 5.95, color=LF_D)
alabel(s, 3.15, 5.95, "クラウド", w=0.9, color=LF_D)

prs.save(OUT)
