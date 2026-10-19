import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

OUT = sys.argv[1]
FONT = "Hiragino Sans"
MONO = "Menlo"
NAVY = RGBColor(0x24, 0x40, 0x4D)
SUB = RGBColor(0x5E, 0x76, 0x82)
GRAY_BG = RGBColor(0xF1, 0xF3, 0xF4)
GRAY_LN = RGBColor(0xC2, 0xCB, 0xD0)
TEAL_BG = RGBColor(0xE6, 0xF5, 0xF2)
TEAL = RGBColor(0x1A, 0xA6, 0x9A)
TEAL_D = RGBColor(0x0E, 0x7C, 0x72)
RED = RGBColor(0xD9, 0x43, 0x3B)
RED_BG = RGBColor(0xFB, 0xEB, 0xEA)
ORANGE = RGBColor(0xF0, 0x8A, 0x24)
ORANGE_BG = RGBColor(0xFD, 0xF1, 0xE4)
ORANGE_D = RGBColor(0xB8, 0x62, 0x0F)
BLUE = RGBColor(0x3B, 0x6E, 0xC9)
BLUE_BG = RGBColor(0xE9, 0xF0, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

IN = 914400
prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]


def inch(v):
    return Emu(int(v * IN))


def text(slide, x, y, w, h, s, size, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         font=FONT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = s.split("\n") if isinstance(s, str) else s
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = line if isinstance(line, list) else [(line, {})]
        for seg, opt in segs:
            r = para.add_run()
            r.text = seg
            r.font.name = opt.get("font", font)
            r.font.size = Pt(opt.get("size", size))
            r.font.bold = opt.get("bold", bold)
            r.font.color.rgb = opt.get("color", color)
            rpr = r._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = etree.SubElement(rpr, qn(tag))
                el.set("typeface", opt.get("font", font))
    return tb


def ctext(slide, x, y, w, h, s, size, color=NAVY, bold=False, font=FONT):
    return text(slide, x, y, w, h, s, size, color, bold, PP_ALIGN.CENTER, font, MSO_ANCHOR.MIDDLE)


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.5,
        radius=0.08, dash=False):
    sh = slide.shapes.add_shape(shape, inch(x), inch(y), inch(w), inch(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = MSO_LINE.DASH
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def arrow(slide, x1, y1, x2, y2, color=SUB, w=2.0, both=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    ends = ["a:tailEnd"] + (["a:headEnd"] if both else [])
    for tag in ends:
        e = etree.SubElement(ln, qn(tag))
        e.set("type", "triangle")
        e.set("w", "med")
        e.set("len", "med")
    return c


def line(slide, x1, y1, x2, y2, color=GRAY_LN, w=1.25, dash=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    return c


def chip(slide, x, y, w, h, s, fill, ln, tc, size=12, bold=True, font=FONT):
    box(slide, x, y, w, h, fill, ln, lw=1.0, radius=0.2)
    ctext(slide, x, y, w, h, s, size, tc, bold, font)


def center_vertically(slide):
    shapes = list(slide.shapes)
    top = min(sh.top for sh in shapes)
    bottom = max(sh.top + sh.height for sh in shapes)
    dy = (prs.slide_height - (bottom - top)) // 2 - top
    for sh in shapes:
        if hasattr(sh, "begin_y"):
            sh.begin_y = sh.begin_y + dy
            sh.end_y = sh.end_y + dy
        else:
            sh.top = sh.top + dy


# ---------- 図1: MCPの全体像（ホスト・クライアント・サーバ） ----------
s = prs.slides.add_slide(BLANK)

# ユーザー
box(s, 0.35, 2.95, 1.35, 1.3, GRAY_BG, GRAY_LN)
ctext(s, 0.35, 2.95, 1.35, 1.3, "ユーザー", 15, NAVY, True)

# ホスト
box(s, 2.45, 0.6, 4.55, 5.0, BLUE_BG, BLUE, radius=0.04)
text(s, 2.7, 0.75, 4.1, 0.4, "MCPホスト（AIアプリ）", 16, BLUE, True)
text(s, 2.7, 1.17, 4.1, 0.3, "例：Claude Desktop、Claude Code", 12, SUB)

box(s, 2.75, 2.75, 1.35, 1.7, WHITE, BLUE)
ctext(s, 2.75, 2.75, 1.35, 1.7, "LLM", 17, BLUE, True)

cy = [1.75, 3.15, 4.55]  # クライアントの上端
cw, ch, cx = 1.85, 0.8, 4.75
for i, y in enumerate(cy):
    box(s, cx, y, cw, ch, WHITE, BLUE)
    ctext(s, cx, y, cw, ch, f"MCPクライアント{i + 1}", 13, NAVY, True)
    arrow(s, 4.1, 3.6, cx, y + ch / 2, BLUE, 1.5, both=True)

# ユーザー ⇄ LLM
arrow(s, 1.7, 3.4, 2.75, 3.4, SUB)
text(s, 1.7, 3.05, 1.05, 0.3, "依頼", 12, SUB, align=PP_ALIGN.CENTER)
arrow(s, 2.75, 3.95, 1.7, 3.95, SUB)
text(s, 1.7, 3.6, 1.05, 0.3, "回答", 12, SUB, align=PP_ALIGN.CENTER)

# サーバ
sx, sw, sh_ = 8.85, 2.2, 0.95
servers = [
    ("MCPサーバA", "ファイル操作", TEAL_BG, TEAL, TEAL_D),
    ("MCPサーバB", "データベース", TEAL_BG, TEAL, TEAL_D),
    ("MCPサーバC", "外部サービス", ORANGE_BG, ORANGE, ORANGE_D),
]
transports = ["stdio", "stdio", "Streamable HTTP"]
targets = ["ファイル", "DB", "Web API"]
# 同じPC / ネット越しの枠
box(s, 8.65, 1.4, 2.6, 2.95, None, TEAL, lw=1.25, radius=0.05, dash=True)
text(s, 8.8, 1.47, 2.4, 0.3, "同じPCで動く", 12, TEAL_D, True)
box(s, 8.65, 4.45, 2.6, 1.4, None, ORANGE, lw=1.25, radius=0.08, dash=True)
text(s, 8.8, 4.5, 2.4, 0.3, "ネット越し", 12, ORANGE_D, True)

for i, ((name, role, bg, ln, tc), tr, tg) in enumerate(zip(servers, transports, targets)):
    y = cy[i] + ch / 2 - sh_ / 2 + (0.2 if i < 2 else 0.25)
    mid = y + sh_ / 2
    box(s, sx, y, sw, sh_, bg, ln)
    text(s, sx, y + 0.13, sw, 0.35, name, 14, tc, True, PP_ALIGN.CENTER)
    text(s, sx, y + 0.52, sw, 0.3, role, 12, NAVY, align=PP_ALIGN.CENTER)
    # クライアント ⇄ サーバ（1対1）
    arrow(s, cx + cw, cy[i] + ch / 2, sx, mid, ln, 2.0, both=True)
    text(s, cx + cw, (cy[i] + ch / 2 + mid) / 2 - 0.42, sx - cx - cw - 0.2, 0.3, tr, 12, tc, True,
         PP_ALIGN.CENTER, MONO if tr == "stdio" else FONT)
    # サーバの先
    box(s, 11.65, mid - 0.35, 1.35, 0.7, GRAY_BG, GRAY_LN)
    ctext(s, 11.65, mid - 0.35, 1.35, 0.7, tg, 13, NAVY, True)
    arrow(s, sx + sw, mid, 11.65, mid, SUB, 1.5, both=True)

# JSON-RPC
text(s, 6.6, 5.75, 2.0, 0.3, "中身はJSON-RPC 2.0", 12, SUB, align=PP_ALIGN.CENTER)

# サーバが提供するもの
text(s, 2.45, 6.12, 3.0, 0.4, "サーバが提供するもの", 13, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
for i, (t, d) in enumerate([("ツール", "LLMが呼ぶ操作"), ("リソース", "読ませるデータ"),
                            ("プロンプト", "指示のひな形")]):
    x = 5.0 + i * 2.75
    chip(s, x, 6.1, 1.15, 0.45, t, TEAL_BG, TEAL, TEAL_D, 12)
    text(s, x + 1.25, 6.1, 1.45, 0.45, d, 12, NAVY, anchor=MSO_ANCHOR.MIDDLE)


# ---------- 図2: ツールを呼ぶまでの流れ（シーケンス） ----------
s = prs.slides.add_slide(BLANK)
lanes = [
    ("ユーザー", GRAY_BG, GRAY_LN, NAVY),
    ("LLM", BLUE_BG, BLUE, BLUE),
    ("MCPホスト\n（クライアント）", BLUE_BG, BLUE, BLUE),
    ("MCPサーバ", TEAL_BG, TEAL, TEAL_D),
    ("ファイル・API など", GRAY_BG, GRAY_LN, NAVY),
]
lx = [1.25, 3.85, 6.45, 9.05, 11.65]
LW, LH, TOP = 2.05, 0.8, 0.3
for (name, bg, ln, tc), x in zip(lanes, lx):
    box(s, x - LW / 2, TOP, LW, LH, bg, ln)
    ctext(s, x - LW / 2, TOP, LW, LH, name, 13, tc, True)
    line(s, x, TOP + LH, x, 7.15)

U, L, H, S, T = range(5)
msgs = [
    (H, S, "① tools/list", False, TEAL_D, True),
    (S, H, "② ツールの一覧", True, TEAL_D, False),
    (U, L, "③ 依頼", False, SUB, False),
    (L, H, "④ ツール名と引数", False, BLUE, False),
    (H, U, "⑤ 実行してよい？", False, RED, False),
    (U, H, "⑥ 許可", True, RED, False),
    (H, S, "⑦ tools/call", False, TEAL_D, True),
    (S, T, "⑧ 実際の処理", False, SUB, False),
    (S, H, "⑨ 結果", True, TEAL_D, False),
    (H, L, "⑩ 結果を渡す", True, BLUE, False),
    (L, U, "⑪ 回答", True, SUB, False),
]
y0, step = 1.6, 0.5
for i, (a, b, label, ret, color, mono) in enumerate(msgs):
    y = y0 + i * step + (0.3 if i >= 2 else 0)
    x1, x2 = lx[a], lx[b]
    pad = 0.06 if x2 > x1 else -0.06
    arrow(s, x1 + pad, y, x2 - pad, y, color, 1.75, dash=ret)
    left, right = min(x1, x2), max(x1, x2)
    if mono:
        lab = [[(label[:2], {}), (label[2:], {"font": MONO})]]
    else:
        lab = label
    text(s, left, y - 0.36, right - left, 0.3, lab, 12, color, True, PP_ALIGN.CENTER)

# 接続時と会話中の区切り
ysep = y0 + 1 * step + 0.4
line(s, 0.2, ysep, 12.95, ysep, GRAY_LN, 1.0)
text(s, 0.2, y0 - 0.38, 1.0, 0.3, "接続時", 12, SUB, True)
text(s, 0.2, ysep + 0.08, 1.0, 0.3, "会話中", 12, SUB, True)


# ---------- 図3: PaperLens の MCP 構成 ----------
s = prs.slides.add_slide(BLANK)

# 左：ホスト2つ
hosts = [("Claude Code", "claude mcp add -s user"),
         ("Claude Desktop", "claude_desktop_config.json")]
hy = [0.85, 3.65]
HW, HH = 3.6, 2.2
for (name, reg), y in zip(hosts, hy):
    box(s, 0.35, y, HW, HH, BLUE_BG, BLUE, radius=0.05)
    text(s, 0.6, y + 0.18, 3.1, 0.4, name, 16, BLUE, True)
    text(s, 0.6, y + 0.68, 3.1, 0.3, "登録", 12, SUB, True)
    text(s, 0.6, y + 0.98, 3.2, 0.3, reg, 12, NAVY, font=MONO)
    chip(s, 0.6, y + 1.45, 2.3, 0.5, "MCPクライアント", WHITE, BLUE, BLUE, 12)

# 中央：PaperLens MCPサーバ
PX, PY, PW, PH = 5.55, 0.85, 4.1, 5.0
box(s, PX, PY, PW, PH, TEAL_BG, TEAL, radius=0.04)
text(s, PX + 0.25, PY + 0.18, PW - 0.5, 0.4, "PaperLens MCPサーバ", 16, TEAL_D, True)
text(s, PX + 0.25, PY + 0.62, PW - 0.5, 0.3, "Node.js / TypeScript SDK", 12, SUB)
text(s, PX + 0.25, PY + 0.92, PW - 0.5, 0.3, "dist-server/mcp.js", 12, NAVY, font=MONO)

text(s, PX + 0.25, PY + 1.45, PW - 0.5, 0.3, "公開するツール", 13, NAVY, True)
tools = [("list_papers", "読み取り", TEAL_BG, TEAL, TEAL_D, MONO),
         ("訳の書き込み", "書き込み", RED_BG, RED, RED, FONT)]
for i, (nm, kind, bg, ln, tc, f) in enumerate(tools):
    y = PY + 1.85 + i * 0.68
    box(s, PX + 0.25, y, 2.45, 0.52, WHITE, ln, lw=1.0, radius=0.15)
    ctext(s, PX + 0.25, y, 2.45, 0.52, nm, 13, NAVY, True, f)
    chip(s, PX + 2.85, y + 0.06, 0.95, 0.4, kind, bg, ln, tc, 12)

# 読み取り専用モード
ry = PY + 3.45
box(s, PX + 0.25, ry, PW - 0.5, 1.3, WHITE, RED, lw=1.25, radius=0.08, dash=True)
text(s, PX + 0.45, ry + 0.15, PW - 0.9, 0.3, "PAPERLENS_MCP_READONLY", 12, RED, True, font=MONO)
text(s, PX + 0.45, ry + 0.55, PW - 0.9, 0.6, "付けると書き込み系の\nツールを公開しない", 12, NAVY)

# ホスト → サーバ（stdio）
for y in hy:
    yc = y + 1.7
    tx = PX
    ty = PY + 1.0 if y < 2 else PY + 4.0
    arrow(s, 0.35 + HW, yc, tx, ty, TEAL_D, 2.0, both=True)
    text(s, 0.35 + HW, (yc + ty) / 2 - 0.42, tx - 0.35 - HW, 0.3, "stdio", 12, TEAL_D, True,
         PP_ALIGN.CENTER, MONO)

# 右：データとアプリ本体
DX, DW = 10.55, 2.45
box(s, DX, 1.75, DW, 1.35, GRAY_BG, GRAY_LN)
text(s, DX, 1.95, DW, 0.35, "論文データ", 14, NAVY, True, PP_ALIGN.CENTER)
text(s, DX, 2.45, DW, 0.3, "PAPERLENS_DATA", 12, SUB, align=PP_ALIGN.CENTER, font=MONO)

box(s, DX, 4.35, DW, 1.35, ORANGE_BG, ORANGE)
text(s, DX, 4.55, DW, 0.35, "PaperLens アプリ", 14, ORANGE_D, True, PP_ALIGN.CENTER)
text(s, DX, 5.05, DW, 0.3, "人が論文を読む画面", 12, NAVY, align=PP_ALIGN.CENTER)

arrow(s, PX + PW, 2.425, DX, 2.425, SUB, 2.0, both=True)
text(s, PX + PW, 2.05, DX - PX - PW, 0.3, "読む・書く", 12, SUB, True, PP_ALIGN.CENTER)
arrow(s, DX + DW / 2, 4.35, DX + DW / 2, 3.1, ORANGE_D, 2.0, both=True)
text(s, DX + DW / 2 + 0.12, 3.55, 1.0, 0.3, "同じデータ", 12, ORANGE_D, True)

# ユーザー（両方を使う）
box(s, 0.35, 6.4, 12.65, 0.6, GRAY_BG, GRAY_LN)
ctext(s, 0.35, 6.4, 12.65, 0.6, "ユーザー", 14, NAVY, True)
arrow(s, 2.15, 5.85, 2.15, 6.4, SUB, 1.75, both=True)
text(s, 2.3, 5.97, 2.0, 0.3, "AIに頼む", 12, SUB, True)
arrow(s, DX + DW / 2, 5.7, DX + DW / 2, 6.4, SUB, 1.75, both=True)
text(s, DX + DW / 2 + 0.12, 5.9, 1.2, 0.3, "自分で読む", 12, SUB, True)


for sl in prs.slides:
    center_vertically(sl)

prs.save(OUT)
print("saved", OUT)
