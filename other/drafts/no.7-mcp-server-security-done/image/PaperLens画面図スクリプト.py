import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

OUT = sys.argv[1]
FONT = "Hiragino Sans"
MONO = "Menlo"
NAVY = RGBColor(0x24, 0x40, 0x4D)
SUB = RGBColor(0x5E, 0x76, 0x82)
GRAY_BG = RGBColor(0xF1, 0xF3, 0xF4)
GRAY_LN = RGBColor(0xC2, 0xCB, 0xD0)
TEAL_BG = RGBColor(0xE6, 0xF5, 0xF2)
TEAL = RGBColor(0x1A, 0xA6, 0x9A)
TEAL_D = RGBColor(0x0E, 0x7C, 0x72)
RED = RGBColor(0xD9, 0x43, 0x3B)
RED_BG = RGBColor(0xFB, 0xEB, 0xEA)
ORANGE = RGBColor(0xF0, 0x8A, 0x24)
ORANGE_BG = RGBColor(0xFD, 0xF1, 0xE4)
ORANGE_D = RGBColor(0xB8, 0x62, 0x0F)
BLUE = RGBColor(0x3B, 0x6E, 0xC9)
BLUE_BG = RGBColor(0xE9, 0xF0, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

IN = 914400
prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]


def inch(v):
    return Emu(int(v * IN))


def text(slide, x, y, w, h, s, size, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         font=FONT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = s.split("\n") if isinstance(s, str) else s
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = line if isinstance(line, list) else [(line, {})]
        for seg, opt in segs:
            r = para.add_run()
            r.text = seg
            r.font.name = opt.get("font", font)
            r.font.size = Pt(opt.get("size", size))
            r.font.bold = opt.get("bold", bold)
            r.font.color.rgb = opt.get("color", color)
            rpr = r._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = etree.SubElement(rpr, qn(tag))
                el.set("typeface", opt.get("font", font))
    return tb


def ctext(slide, x, y, w, h, s, size, color=NAVY, bold=False, font=FONT):
    return text(slide, x, y, w, h, s, size, color, bold, PP_ALIGN.CENTER, font, MSO_ANCHOR.MIDDLE)


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.5,
        radius=0.08, dash=False):
    sh = slide.shapes.add_shape(shape, inch(x), inch(y), inch(w), inch(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = MSO_LINE.DASH
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def arrow(slide, x1, y1, x2, y2, color=SUB, w=2.0, both=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    ends = ["a:tailEnd"] + (["a:headEnd"] if both else [])
    for tag in ends:
        e = etree.SubElement(ln, qn(tag))
        e.set("type", "triangle")
        e.set("w", "med")
        e.set("len", "med")
    return c


def line(slide, x1, y1, x2, y2, color=GRAY_LN, w=1.25, dash=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    return c


def chip(slide, x, y, w, h, s, fill, ln, tc, size=12, bold=True, font=FONT):
    box(slide, x, y, w, h, fill, ln, lw=1.0, radius=0.2)
    ctext(slide, x, y, w, h, s, size, tc, bold, font)


def center_vertically(slide):
    shapes = list(slide.shapes)
    top = min(sh.top for sh in shapes)
    bottom = max(sh.top + sh.height for sh in shapes)
    dy = (prs.slide_height - (bottom - top)) // 2 - top
    for sh in shapes:
        if hasattr(sh, "begin_y"):
            sh.begin_y = sh.begin_y + dy
            sh.end_y = sh.end_y + dy
        else:
            sh.top = sh.top + dy



# ---------- PaperLens の画面と MCP サーバの関係 ----------
import os
IMG = sys.argv[2]
s = prs.slides.add_slide(BLANK)
PXW, PXH = 1470, 923
X0, Y0, W = 0.3, 0.98, 8.8
K = W / PXW
s.shapes.add_picture(IMG, inch(X0), inch(Y0), inch(W), inch(PXH * K))
box(s, X0, Y0, W, PXH * K, None, GRAY_LN, MSO_SHAPE.RECTANGLE, lw=1.0)


def px(x, y):
    return X0 + x * K, Y0 + y * K


# 赤枠：MCP の書き込み系ツールで書き込まれた部分
rx1, ry1 = px(710, 262)
rx2, ry2 = px(1450, 803)
box(s, rx1, ry1, rx2 - rx1, ry2 - ry1, None, RED, MSO_SHAPE.RECTANGLE, lw=3.0)

# 右：ホスト → MCPサーバ → 画面
CX, CW = 9.9, 3.15
box(s, CX, 1.0, CW, 0.95, BLUE_BG, BLUE)
ctext(s, CX, 1.0, CW, 0.95, "Claude Code\nClaude Desktop", 14, BLUE, True)
arrow(s, CX + CW / 2, 1.95, CX + CW / 2, 2.65, TEAL_D, 2.0)
text(s, CX + CW / 2 + 0.12, 2.15, 1.5, 0.3, "tools/call", 12, TEAL_D, True, font=MONO)

box(s, CX, 2.65, CW, 1.65, TEAL_BG, TEAL)
text(s, CX, 2.8, CW, 0.4, "PaperLens MCPサーバ", 14, TEAL_D, True, PP_ALIGN.CENTER)
chip(s, CX + 0.35, 3.45, CW - 0.7, 0.5, "書き込み系ツール", RED_BG, RED, RED, 13)

ay = 3.7
arrow(s, CX, ay, rx2, ay, RED, 2.5)
text(s, rx2, ay - 0.36, CX - rx2, 0.3, "書き込む", 12, RED, True, PP_ALIGN.CENTER)

box(s, CX, 4.6, CW, 1.15, None, RED, lw=1.25, radius=0.08, dash=True)
text(s, CX + 0.2, 4.72, CW - 0.4, 0.3, "PAPERLENS_MCP_READONLY", 12, RED, True, font=MONO)
text(s, CX + 0.2, 5.1, CW - 0.4, 0.5, "付けると公開されない", 12, NAVY)

prs.save(OUT)
print("saved", OUT)
