# no.13 Agentic Internet の図を作る（python-pptx）
# 使い方: python3 図生成スクリプト.py <出力pptxのパス>
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

NAVY = RGBColor(0x2E, 0x5E, 0x8E)
TEAL = RGBColor(0x1E, 0x8A, 0x82)
RED = RGBColor(0xC0, 0x39, 0x2B)
TEXT = RGBColor(0x1F, 0x3A, 0x5C)
SUB = RGBColor(0x6B, 0x72, 0x80)
SLATE = RGBColor(0x2B, 0x3A, 0x4A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_BG = RGBColor(0xEE, 0xF3, 0xF9)
TEAL_BG = RGBColor(0xE8, 0xF5, 0xF3)
TEAL_MID = RGBColor(0xC9, 0xE8, 0xE4)
GRAY_BG = RGBColor(0xF3, 0xF4, 0xF6)
GRAY_MID = RGBColor(0xE5, 0xE7, 0xEB)
FONT = "Noto Sans CJK JP"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(8.5)


def new_slide():
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_text(shape, lines, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER):
    """lines: [(text, size, bold, color), ...]"""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.08))
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, (t, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
        rpr = r._r.get_or_add_rPr()
        ea = rpr.find(qn("a:ea"))
        if ea is None:
            ea = etree.SubElement(rpr, qn("a:ea"))
        ea.set("typeface", FONT)


def rrect(s, x, y, w, h, fill, line, lw=1.4, adj=0.16, dash=False):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = adj
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = 4  # dash
    return sh


def rect(s, x, y, w, h, fill, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.4)
    return sh


def node(s, x, y, w, h, title, sub=None, border=NAVY, fill=WHITE, tsize=18, ssize=16, tcolor=TEXT):
    sh = rrect(s, x, y, w, h, fill, border)
    lines = [(title, tsize, True, tcolor)]
    if sub:
        lines.append((sub, ssize, False, SUB))
    set_text(sh, lines)
    return sh


def zone(s, x, y, w, h, label, accent, bg, tabw=None):
    rrect(s, x, y, w, h, bg, accent, lw=1.6, adj=0.05)
    tw = tabw or (0.6 + 0.3 * len(label))
    tab = rrect(s, x + 0.3, y - 0.26, tw, 0.52, accent, None, adj=0.5)
    set_text(tab, [(label, 18, True, WHITE)])


def label(s, x, y, w, h, text, size=16, color=SUB, bold=False, align=PP_ALIGN.CENTER):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(tb, [(text, size, bold, color)], align=align)
    return tb


def arrow(s, x1, y1, x2, y2, color=SUB, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    if dash:
        c.line.dash_style = 4
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c

def card(s, x, y, w, h, lines, border=NAVY, fill=WHITE):
    """lines: [(text, size, bold, color), ...]"""
    sh = rrect(s, x, y, w, h, fill, border)
    set_text(sh, lines)
    return sh


def line(s, x1, y1, x2, y2, color=SUB):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    return c


RED_BG = RGBColor(0xFB, 0xEC, 0xEA)

# ---------- 図1: Cloudflareで見えている通信の量 ----------
# 出典: Cloudflare Blog「The Internet has a second audience」(2026-09-30)
s = new_slide()
BASE = 7.2
SCALE = 5.2 / 150  # 1億5,000万を5.2インチに

zone(s, 0.5, 0.75, 7.2, 7.25, "HTTPリクエスト数（1秒あたり）", NAVY, NAVY_BG, tabw=5.4)
h1, h2, hp = 63 * SCALE, 115 * SCALE, 150 * SCALE
rect(s, 1.5, BASE - h1, 2.0, h1, NAVY)
rect(s, 4.6, BASE - h2, 2.0, h2, NAVY)
rrect(s, 4.6, BASE - hp, 2.0, hp - h2, None, NAVY, adj=0.0, dash=True)
label(s, 1.2, BASE - h1 - 0.6, 2.6, 0.5, "6,300万", 20, TEXT, True)
label(s, 4.4, BASE - h2 + 0.1, 2.4, 0.5, "1億1,500万", 20, WHITE, True)
label(s, 3.9, BASE - hp - 0.6, 3.4, 0.5, "ピーク時 1億5,000万超", 16, NAVY, True)
arrow(s, 3.6, BASE - h1 + 0.1, 4.5, BASE - h2 + 0.5, NAVY)
label(s, 2.6, BASE - h2 - 0.15, 1.9, 0.5, "約1.8倍", 20, NAVY, True)
line(s, 0.9, BASE, 7.3, BASE)
label(s, 1.3, BASE + 0.08, 2.4, 0.45, "2024年末", 16, TEXT)
label(s, 4.4, BASE + 0.08, 2.4, 0.45, "2026年9月", 16, TEXT)

zone(s, 8.05, 0.75, 4.8, 7.25, "中身の変化", TEAL, TEAL_BG)
card(s, 8.3, 1.4, 4.3, 2.2, [
    ("18倍以上", 30, True, TEAL),
    ("AIエージェントの", 16, False, TEXT),
    ("1日あたりリクエスト（1年間）", 16, False, TEXT),
], border=TEAL)
card(s, 8.3, 3.85, 4.3, 1.95, [
    ("最大4割減", 30, True, RED),
    ("人からのトラフィック", 16, False, TEXT),
    ("（業界によっては1年未満で）", 16, False, SUB),
], border=RED)
card(s, 8.3, 6.05, 4.3, 1.7, [
    ("週5,000億超", 28, True, NAVY),
    ("身元が確認されたボットの", 16, False, TEXT),
    ("リクエスト", 16, False, TEXT),
], border=NAVY)


# ---------- 図2: アイソレートで何が変わるか ----------
s = new_slide()
NW, NH = 2.2, 1.5
XS = [3.05, 5.6, 8.15, 10.7]

start = rrect(s, 0.35, 3.1, 1.95, 1.5, SLATE, None)
set_text(start, [("エージェント", 18, True, WHITE), ("コードを書いた", 16, False, WHITE)])

# 上：従来（コンテナで隔離）
zone(s, 2.8, 0.65, 10.3, 3.0, "従来：コンテナで隔離", NAVY, NAVY_BG)
UY = 1.35
rrect(s, 2.95, 1.15, 4.95, 1.9, None, RED, dash=True, adj=0.1)
node(s, XS[0], UY, NW, NH, "コンテナを作る", "OSの環境を用意", border=NAVY)
node(s, XS[1], UY, NW, NH, "ランタイム起動", "Node.js など", border=NAVY)
node(s, XS[2], UY, NW, NH, "コードを実行", "コンテナの中で", border=NAVY)
card(s, XS[3], UY, NW, NH, [
    ("数百ミリ秒", 18, True, RED), ("メモリ 数百MB", 15, False, TEXT), ("同時数に上限", 15, False, TEXT)
], border=RED, fill=RED_BG)
label(s, 2.95, 3.1, 4.95, 0.45, "依頼のたびに毎回やる準備", 16, RED, True)

# 下：アイソレート（Dynamic Workers）
zone(s, 2.8, 4.1, 10.3, 3.15, "アイソレートで隔離（Dynamic Workers）", TEAL, TEAL_BG, tabw=6.4)
LY = 4.8
rrect(s, 2.95, 4.6, 7.5, 1.9, None, TEAL, dash=True, adj=0.1)
node(s, XS[0], LY, NW, NH, "親Worker", "エージェント本体", border=TEAL)
node(s, XS[1], LY, NW, NH, "アイソレート", "LOADER.load()", border=TEAL)
node(s, XS[2], LY, NW, NH, "コードを実行", "権限は親が制限", border=TEAL)
card(s, XS[3], LY, NW, NH, [
    ("数ミリ秒", 18, True, TEAL), ("メモリ 数MB", 15, False, TEXT), ("依頼ごとに作成", 15, False, TEXT)
], border=TEAL, fill=TEAL_MID)
label(s, 2.95, 6.55, 7.5, 0.45, "すべて同じプロセスの中（ランタイムの準備は1回だけ）", 16, TEAL, True)

for y in (UY, LY):
    for i in range(3):
        arrow(s, XS[i] + NW, y + NH / 2, XS[i + 1], y + NH / 2, SUB)
arrow(s, 2.3, 3.55, XS[0], UY + NH / 2, SUB)
arrow(s, 2.3, 4.15, XS[0], LY + NH / 2, SUB)


# ---------- 図3: コンテナとアイソレートの構成の比較 ----------
s = new_slide()
CW, GAP = 1.7, 0.15


def layer(x, y, w, text, fill=GRAY_MID, color=TEXT):
    b = rrect(s, x, y, w, 0.55, fill, None, adj=0.2)
    set_text(b, [(text, 16, True, color)])


def column(cx, top, title, boxes, accent):
    h = 0.45 + 0.7 * len(boxes) + 0.15
    rrect(s, cx, top, CW, h, None, accent, adj=0.08)
    label(s, cx, top + 0.03, CW, 0.42, title, 16, accent, True)
    for j, (text, fill) in enumerate(boxes):
        b = rrect(s, cx + 0.1, top + 0.5 + 0.7 * j, CW - 0.2, 0.6, fill, accent if fill == WHITE else None, adj=0.2)
        set_text(b, [(text, 15, True, TEXT)])


# 左：コンテナ
zone(s, 0.4, 0.65, 6.0, 5.6, "コンテナ", NAVY, NAVY_BG)
for i in range(3):
    column(0.7 + i * (CW + GAP), 1.25, f"コンテナ{i + 1}",
           [("コード", WHITE), ("ランタイム", GRAY_MID), ("ライブラリ等", GRAY_MID)], NAVY)
layer(0.7, 4.15, 5.4, "コンテナランタイム（containerd など）")
layer(0.7, 4.8, 5.4, "ホストOS（Linuxカーネル）")
layer(0.7, 5.45, 5.4, "ハードウェア")

# 右：アイソレート
zone(s, 6.95, 0.65, 6.0, 5.6, "アイソレート", TEAL, TEAL_BG)
for i in range(3):
    column(7.25 + i * (CW + GAP), 1.25, f"アイソレート{i + 1}",
           [("コード", WHITE), ("専用のヒープ", TEAL_MID)], TEAL)
layer(7.25, 3.5, 5.4, "V8エンジン（V8サンドボックス）", TEAL, WHITE)
layer(7.25, 4.15, 5.4, "Workersのランタイム（workerd）", TEAL, WHITE)
layer(7.25, 4.8, 5.4, "ホストOS（Linuxカーネル）")
layer(7.25, 5.45, 5.4, "ハードウェア")

# 下：数字と隔離のしくみ
node(s, 0.4, 6.5, 2.9, 1.05, "数百ミリ秒", "起動にかかる時間", border=NAVY, tsize=20)
node(s, 3.5, 6.5, 2.9, 1.05, "数百MB", "動かすのに要るメモリ", border=NAVY, tsize=20)
node(s, 6.95, 6.5, 2.9, 1.05, "数ミリ秒", "起動にかかる時間", border=TEAL, tsize=20, tcolor=TEAL)
node(s, 10.05, 6.5, 2.9, 1.05, "数MB", "動かすのに要るメモリ", border=TEAL, tsize=20, tcolor=TEAL)
card(s, 0.4, 7.7, 6.0, 0.6, [("隔離：カーネルの名前空間・cgroups", 16, True, NAVY)], border=NAVY)
card(s, 6.95, 7.7, 6.0, 0.6, [("隔離：V8サンドボックス＋2層目のサンドボックス", 16, True, TEAL)], border=TEAL)

prs.save(sys.argv[1])
